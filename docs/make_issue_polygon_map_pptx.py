"""
Generate tools/rcn-issue-polygon-map-intro.pptx — the Issue Polygon Map intro deck.

Companion to tools/issue-polygon-map-intro.html; same argument, deck shape.
Regenerate after editing:  python3 docs/make_issue_polygon_map_pptx.py
Requires: pip install python-pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette — gold/amber for the spatial tool, matching the intro page ────────
GD     = RGBColor(0xd4, 0xa0, 0x17)   # issue-polygon gold
GD_LT  = RGBColor(0xfe, 0xf3, 0xc7)   # light gold
GD_BG  = RGBColor(0xff, 0xfb, 0xeb)   # near-white gold
GD_DK  = RGBColor(0x92, 0x40, 0x0e)   # dark amber
B      = RGBColor(0x1a, 0x56, 0xa4)   # blue — state jurisdiction
B_LT   = RGBColor(0xdb, 0xea, 0xfe)
PU     = RGBColor(0x6b, 0x21, 0xa8)   # purple — county
PU_LT  = RGBColor(0xf3, 0xe8, 0xff)
OR     = RGBColor(0x9a, 0x34, 0x12)   # orange — municipality
OR_LT  = RGBColor(0xff, 0xed, 0xd5)
GRN    = RGBColor(0x16, 0x65, 0x34)   # pro
GRN_LT = RGBColor(0xdc, 0xfc, 0xe7)
RD     = RGBColor(0x99, 0x1b, 0x1b)   # con
RD_LT  = RGBColor(0xfe, 0xe2, 0xe2)
BLACK  = RGBColor(0x0f, 0x17, 0x2a)
GREY   = RGBColor(0x47, 0x55, 0x69)
LGREY  = RGBColor(0x94, 0xa3, 0xb8)
SLATE  = RGBColor(0xe2, 0xe8, 0xf0)
WHITE  = RGBColor(0xff, 0xff, 0xff)
OFF    = RGBColor(0xf8, 0xfa, 0xfc)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Primitives ───────────────────────────────────────────────────────────────

def slide():
    return prs.slides.add_slide(BLANK)


def bg(sl, color):
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = color


def rect(sl, x, y, w, h, fill=WHITE, line=None, lw=1.5):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def txt(sl, text, x, y, w, h, size=20, bold=False, italic=False,
        color=BLACK, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return tb


def txlines(sl, lines, x, y, w, h, sizes=None, bolds=None, colors=None,
            align=PP_ALIGN.LEFT, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(sizes[i] if isinstance(sizes, list) else (sizes or 18))
        r.font.bold = bolds[i] if isinstance(bolds, list) else (bolds or False)
        r.font.color.rgb = colors[i] if isinstance(colors, list) else (colors or BLACK)
        r.font.name = 'Calibri'
    return tb


def box_txt(sl, lines, x, y, w, h, fill=GD_BG, line=GD, lw=1.5,
            size=15, bold_first=True, text_color=BLACK,
            top_color=None, align=PP_ALIGN.CENTER):
    sh = rect(sl, x, y, w, h, fill=fill, line=line, lw=lw)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)
    for i, line_text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(size if i > 0 else size + 1)
        r.font.bold = (bold_first and i == 0)
        r.font.color.rgb = (top_color or text_color) if i == 0 else text_color
        r.font.name = 'Calibri'
    return sh


def label(sl, text, color=GD_DK):
    txt(sl, text, 0.45, 0.18, 8, 0.3, size=9, bold=True, color=color)


def bar(sl, color=GD, y=0.55, h=0.035):
    rect(sl, 0, y, 13.33, h, fill=color)


def title_txt(sl, text, color=BLACK, size=34, y=0.62):
    txt(sl, text, 0.45, y, 12.4, 1.3, size=size, bold=True, color=color)


def foot(sl, text):
    txt(sl, text, 0.45, 6.95, 12.4, 0.35, size=10, color=LGREY)


# ── 1 · Title ────────────────────────────────────────────────────────────────
s = slide()
bg(s, BLACK)
rect(s, 0, 0, 13.33, 0.14, fill=GD)
txt(s, "RCN TOOLSET  ·  SPATIAL LAYER", 0.9, 1.75, 11, 0.4,
    size=13, bold=True, color=GD)
txt(s, "Issue Polygon Map", 0.9, 2.25, 11.5, 1.25, size=54, bold=True, color=WHITE)
txlines(s, [
    "Draw the boundary of a civic issue.",
    "See which government's rule lands inside it.",
    "Record where every property owner actually stands.",
], 0.9, 3.7, 11, 1.6, sizes=22, colors=SLATE, spacing=8)
txt(s, "Part of the RCN family alongside the Graph Tool (cause) and the Timeline (time)",
    0.9, 6.3, 11, 0.4, size=13, italic=True, color=LGREY)

# ── 2 · The problem ──────────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "THE PROBLEM")
bar(s)
title_txt(s, "Most civic arguments are arguments about a boundary")
txlines(s, [
    "Who is inside the rule and who is outside it.",
    "Which government's law applies here and not there.",
    "How many of the people actually affected want the change.",
], 0.55, 2.1, 12.2, 1.8, sizes=24, colors=GREY, spacing=12)
box_txt(s, [
    "And almost nobody draws it.",
    "The boundary stays an assumption everyone believes they share — until a hearing, when it turns out they never did.",
], 0.55, 4.35, 12.2, 1.5, fill=GD_BG, line=GD, size=17, top_color=GD_DK)
foot(s, "The map makes the boundary an object you can draw, share and argue with.")

# ── 3 · The third kind of region ─────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHERE IT FITS")
bar(s)
title_txt(s, "Every RCN tool is regions plus typed relations")
box_txt(s, ["🔗  CAUSALITY", "Nodes + causal edges", "The Graph Tool"],
        0.55, 2.35, 3.9, 2.0, fill=WHITE, line=SLATE, size=15, top_color=GREY)
box_txt(s, ["⏳  TIME", "Intervals + before / meets", "The Timeline"],
        4.72, 2.35, 3.9, 2.0, fill=WHITE, line=SLATE, size=15, top_color=GREY)
box_txt(s, ["🗺️  SPACE", "Parcels + jurisdictions", "This tool"],
        8.88, 2.35, 3.9, 2.0, fill=GD_BG, line=GD, lw=2.5, size=15, top_color=GD_DK)
txt(s, "Here a region is a parcel or a boundary — and the relations are the ones geography gives you for free: "
       "inside, outside, overlapping, adjacent.",
    0.55, 4.75, 12.2, 1.0, size=18, color=GREY)
foot(s, "Same shape underneath. Different dimension.")

# ── 4 · An issue, not a map ──────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "THE KEY DISTINCTION")
bar(s)
title_txt(s, "It holds an issue, not a map")
txt(s, "A GIS viewer shows you land. This shows you one question — and the land is there only because the question has a footprint.",
    0.55, 1.95, 12.2, 0.8, size=19, color=GREY)
cols = [
    ("ISSUE POLYGON", "The boundary that defines the scope", GD_LT, GD_DK),
    ("PARCELS", "Properties inside, each with a stance", GRN_LT, GRN),
    ("BOUNDARIES", "The governments that overlap it", PU_LT, PU),
    ("LAWS", "The rules in play, tagged by jurisdiction", B_LT, B),
    ("ACTIONS", "What nobody has done yet", OR_LT, OR),
]
x = 0.55
for name, desc, fill, col in cols:
    box_txt(s, [name, desc], x, 3.0, 2.35, 1.75, fill=fill, line=col, size=13, top_color=col)
    x += 2.46
box_txt(s, [
    "Everything except stance is public record.",
    "That is what makes this an instrument you can build before you have spoken to anybody.",
], 0.55, 5.15, 12.2, 1.3, fill=WHITE, line=SLATE, size=16, top_color=BLACK)

# ── 5 · The jurisdiction stack ───────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHY IT MATTERS")
bar(s)
title_txt(s, "The layers are governments, not map furniture")
txt(s, "Superior, Arizona. The question isn't “are chickens allowed.” It's whose rule wins.",
    0.55, 1.95, 12.2, 0.5, size=18, italic=True, color=GREY)
rows = [
    ("STATE", "HB2325 (2024) limits how far a town can restrict residential fowl — but only for code adopted after it passed.", B_LT, B),
    ("COUNTY", "Pinal County's animal ordinance sets a concrete standard: max 6 hens, 100 ft setback.", PU_LT, PU),
    ("MUNICIPALITY", "Town code may prohibit fowl outright — and may predate the preemption. Nobody has verified which.", OR_LT, OR),
]
y = 2.7
for tag, desc, fill, col in rows:
    rect(s, 0.55, y, 2.3, 0.95, fill=fill, line=col, lw=1.5)
    txt(s, tag, 0.6, y + 0.28, 2.2, 0.4, size=14, bold=True, color=col, align=PP_ALIGN.CENTER)
    rect(s, 2.95, y, 9.83, 0.95, fill=WHITE, line=SLATE, lw=1.2)
    txt(s, desc, 3.15, y + 0.2, 9.5, 0.7, size=14, color=GREY)
    y += 1.12
box_txt(s, ["Three governments. Three answers. One back yard."],
        0.55, 6.15, 12.2, 0.6, fill=GD_BG, line=GD, size=17, top_color=GD_DK)

# ── 6 · Stance, honestly ─────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "THE HONEST BIT")
bar(s)
title_txt(s, "Every parcel is Unknown until somebody asks")
box_txt(s, ["PRO", "They said so"], 0.55, 2.3, 3.9, 1.3, fill=GRN_LT, line=GRN, size=16, top_color=GRN)
box_txt(s, ["CON", "They said so"], 4.72, 2.3, 3.9, 1.3, fill=RD_LT, line=RD, size=16, top_color=RD)
box_txt(s, ["UNKNOWN", "Nobody has asked"], 8.88, 2.3, 3.9, 1.3, fill=SLATE, line=GREY, size=16, top_color=GREY)
txlines(s, [
    "Public record gives you lot size, zoning and ownership. It does not give you what a household thinks.",
    "Inferring a stance from a yard sign or a neighbour's opinion puts words in someone's mouth — on a map you intend to show a council.",
], 0.55, 3.95, 12.2, 1.5, sizes=17, colors=GREY, spacing=10)
box_txt(s, [
    "A large Unknown count is not a failure. It is the work remaining.",
    "And it is the number to bring when someone claims the neighbourhood is against something.",
], 0.55, 5.35, 12.2, 1.35, fill=GD_BG, line=GD, size=16, top_color=GD_DK)

# ── 7 · What you get ─────────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHAT IT DOES")
bar(s)
title_txt(s, "What you get")
items = [
    ("Live tally", "Pro / con / unknown, updating as you work"),
    ("Deep links", "A URL that opens on one parcel, popup already open"),
    ("Drawn polygons", "Sketch a boundary, name it, drag its vertices to reshape"),
    ("Basemaps", "Satellite for arguments about setbacks and what's really built"),
    ("GeoJSON round-trip", "Export the whole issue, reload it intact, or open it in QGIS"),
    ("SVG · PNG · Print", "For a flyer, a council packet, or a door-knock sheet"),
]
x, y = 0.55, 2.2
for i, (name, desc) in enumerate(items):
    box_txt(s, [name, desc], x, y, 3.9, 1.35, fill=WHITE, line=SLATE,
            size=13, top_color=GD_DK, align=PP_ALIGN.LEFT)
    x += 4.16
    if i % 3 == 2:
        x = 0.55
        y += 1.55
foot(s, "Field-level schema and its traps: tools/schemas/issue-polygon-map.md")

# ── 8 · Sequence ─────────────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHEN TO USE IT")
bar(s)
title_txt(s, "It needs nobody's permission")
txt(s, "The RCN instruments are ordered by how much relationship they require. This one is near the front.",
    0.55, 1.95, 12.2, 0.5, size=18, color=GREY)
box_txt(s, [
    "NO CONSENT NEEDED",
    "Issue Polygon Map · Graph Tool · EIP sketch · FedWiki pages",
    "Assembled from public record. Build it before the first conversation.",
], 0.55, 2.75, 6.0, 2.1, fill=GD_BG, line=GD, lw=2.5, size=14, top_color=GD_DK)
box_txt(s, [
    "RELATIONSHIP FIRST",
    "e-VSM Survey · CAM field test · SODOTO · CfA-dSC",
    "These ask something of people. They come later.",
], 6.78, 2.75, 6.0, 2.1, fill=WHITE, line=SLATE, size=14, top_color=GREY)
box_txt(s, [
    "A map built in an afternoon is a far better opening than an opinion.",
], 0.55, 5.25, 12.23, 0.75, fill=WHITE, line=SLATE, size=17, top_color=BLACK)
foot(s, "What it can't do alone: say why the situation persists (Graph Tool), or when anything happened (Timeline).")

# ── 9 · Close ────────────────────────────────────────────────────────────────
s = slide()
bg(s, BLACK)
rect(s, 0, 0, 13.33, 0.14, fill=GD)
txt(s, "THE BET UNDERNEATH", 0.9, 1.5, 11, 0.4, size=13, bold=True, color=GD)
txlines(s, [
    "A boundary drawn in public,",
    "with its rules and its unknowns attached,",
    "changes the argument more than another opinion does.",
], 0.9, 2.2, 11.5, 2.4, sizes=32, bolds=True, colors=WHITE, spacing=12)
txt(s, "tools/issue-polygon-map.html", 0.9, 5.3, 11, 0.5, size=18, color=GD)
txt(s, "Introduction · User Manual · Schema reference in tools/schemas/",
    0.9, 5.85, 11, 0.4, size=13, color=LGREY)

# ── Save ─────────────────────────────────────────────────────────────────────
here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, '..', 'tools', 'rcn-issue-polygon-map-intro.pptx')
prs.save(out)
print('Wrote', os.path.normpath(out), '—', len(prs.slides.__iter__.__self__._sldIdLst), 'slides')
