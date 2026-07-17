"""
Generate whatcom-coop-slides-berwick.pptx — "Where Health Actually Lives"
Personal dispatch to Don Berwick · July 2026

Requires: pip install python-pptx
Run:      python3 docs/make_berwick_pptx.py
Output:   docs/whatcom-coop-slides-berwick.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'whatcom-coop-slides-berwick.pptx')

# ── Palette ───────────────────────────────────────────────────────────────────
NAVY       = RGBColor(0x1e, 0x3a, 0x5f)
NAVY_LT    = RGBColor(0x2a, 0x4f, 0x7c)
DARK_TEAL  = RGBColor(0x0f, 0x4c, 0x5c)
TEAL       = RGBColor(0x0f, 0x76, 0x6e)
TEAL_LT    = RGBColor(0xf0, 0xfd, 0xfa)
TEAL_BD    = RGBColor(0x99, 0xf6, 0xe4)
BLUE_LT    = RGBColor(0xef, 0xf6, 0xff)
BLUE_BD    = RGBColor(0x93, 0xc5, 0xfd)
GREEN_LT   = RGBColor(0xf0, 0xfd, 0xf4)
GREEN_BD   = RGBColor(0x86, 0xef, 0xac)
GREEN_DARK = RGBColor(0x16, 0x65, 0x34)
GREEN_TEXT = RGBColor(0x14, 0x53, 0x2d)
GREEN_BOX  = RGBColor(0xdc, 0xfc, 0xe7)
GREEN_BR   = RGBColor(0x22, 0xc5, 0x5e)
PURPLE_LT  = RGBColor(0xfa, 0xf5, 0xff)
PURPLE_BD  = RGBColor(0xc4, 0xb5, 0xfd)
AMBER_BD   = RGBColor(0xfc, 0xd3, 0x4d)
AMBER_LT   = RGBColor(0xff, 0xfb, 0xeb)
WHITE      = RGBColor(0xff, 0xff, 0xff)
BLACK      = RGBColor(0x1a, 0x1a, 0x1a)
GREY_BG    = RGBColor(0xf8, 0xf9, 0xfa)
SURFACE    = RGBColor(0xff, 0xff, 0xff)
TEXT_MED   = RGBColor(0x33, 0x33, 0x33)
TEXT_MUTED = RGBColor(0x55, 0x55, 0x55)
TEXT_GREY  = RGBColor(0x88, 0x88, 0x88)
BORDER     = RGBColor(0xdd, 0xe3, 0xed)
FLOW_GREEN = RGBColor(0xf0, 0xfd, 0xf4)
FLOW_BD_G  = RGBColor(0x22, 0xc5, 0x5e)
DISPATCH   = RGBColor(0xf0, 0xfd, 0xfa)
DISPATCH_BD= RGBColor(0x14, 0xb8, 0xa6)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

# ── Primitives ────────────────────────────────────────────────────────────────
def slide():
    return prs.slides.add_slide(BLANK)

def bg(sl, color):
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = color

def rect(sl, x, y, w, h, fill):
    shp = sl.shapes.add_shape(1, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    return shp

def txt(sl, text, x, y, w, h, size=16, bold=False, color=BLACK,
        align=PP_ALIGN.LEFT, italic=False, wrap=True):
    tb = sl.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb

def label_tag(sl, text, x, y=Inches(0.12)):
    txt(sl, text.upper(), x, y, Inches(10), Inches(0.28),
        size=9, bold=True, color=TEXT_GREY)

def footer(sl):
    txt(sl, 'Where Health Actually Lives · For Don Berwick · July 2026 · CC 4.0',
        Inches(0.6), Inches(7.1), Inches(12), Inches(0.3),
        size=9, color=TEXT_GREY)

def card(sl, x, y, w, h, fill, border):
    shp = sl.shapes.add_shape(1, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(1.5)
    return shp

def card_content(sl, x, y, w, lbl, title, body,
                 lbl_color=TEAL, title_color=NAVY, body_color=TEXT_MED,
                 fill=SURFACE, border=BORDER):
    card(sl, x, y, w, Inches(1.75), fill, border)
    txt(sl, lbl.upper(), x + Inches(0.16), y + Inches(0.12),
        w - Inches(0.32), Inches(0.22), size=9, bold=True, color=lbl_color)
    txt(sl, title, x + Inches(0.16), y + Inches(0.34),
        w - Inches(0.32), Inches(0.4), size=12, bold=True, color=title_color)
    txt(sl, body, x + Inches(0.16), y + Inches(0.74),
        w - Inches(0.32), Inches(0.9), size=10, color=body_color)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 1 — Cover
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, NAVY)

txt(sl, 'FOR DON BERWICK · PERSONAL · JULY 2026',
    Inches(0.7), Inches(1.5), Inches(11.93), Inches(0.35),
    size=11, bold=True, color=RGBColor(0xb0, 0xb8, 0xcc),
    align=PP_ALIGN.CENTER)

txt(sl, 'Where Health\nActually Lives',
    Inches(0.7), Inches(2.0), Inches(11.93), Inches(1.8),
    size=52, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

txt(sl, 'Thirty years of work, a few unexpected obstacles, and a structure\n'
        'that is finally ready — a personal dispatch from Marc Pierson',
    Inches(1.5), Inches(3.9), Inches(10.33), Inches(0.9),
    size=18, color=RGBColor(0xb0, 0xc8, 0xe8), align=PP_ALIGN.CENTER)

txt(sl, 'Whatcom Wealth and Health  ·  My Personal Health Supporter  ·  Open Source CC 4.0',
    Inches(0.7), Inches(5.1), Inches(11.93), Inches(0.3),
    size=10, color=RGBColor(0x6a, 0x8f, 0xb0), align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 2 — The Founding Quote
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, NAVY)

txt(sl, '\u201cMost of the choices and actions that create and preserve\n'
        'individual health occur in the home.\u201d',
    Inches(1.0), Inches(1.9), Inches(11.33), Inches(2.0),
    size=28, italic=True, color=WHITE, align=PP_ALIGN.CENTER)

txt(sl, 'This is not a new idea. It is the idea we have never fully organized around.',
    Inches(1.0), Inches(4.1), Inches(11.33), Inches(0.6),
    size=16, color=RGBColor(0xb0, 0xc8, 0xe8), align=PP_ALIGN.CENTER)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 3 — Neighborhoods Are the Health System
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, SURFACE)

label_tag(sl, 'The organizing insight', Inches(0.7))

txt(sl, 'Neighborhoods are the true health system.\nThey always have been.',
    Inches(0.7), Inches(0.5), Inches(11.5), Inches(1.0),
    size=30, bold=True, color=NAVY)

txt(sl, 'The relationships, habits, knowledge, mutual support, and sense of purpose that exist at '
        'neighborhood scale — these determine the majority of health outcomes. The medical system '
        'responds to illness with skill and dedication. But it does not produce health. Neighborhoods do.',
    Inches(0.7), Inches(1.65), Inches(11.5), Inches(1.1),
    size=15, color=TEXT_MED)

txt(sl, 'What has been missing is not the right clinic, the right technology, or the right payment model. '
        'What has been missing is giving people and their communities everything they need and want to do '
        'what they were already doing \u2014 and to do it far better.',
    Inches(0.7), Inches(2.85), Inches(11.5), Inches(1.1),
    size=15, color=TEXT_MED)

txt(sl, 'That is what we are building.',
    Inches(0.7), Inches(4.05), Inches(11.5), Inches(0.4),
    size=15, bold=True, color=TEAL)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 4 — The Untethered Record
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, GREY_BG)

label_tag(sl, 'My Shared Care Plan \u2014 the founding act', Inches(0.7))

txt(sl, 'The first personal health record\nuntethered from the medical system',
    Inches(0.7), Inches(0.5), Inches(11.5), Inches(1.0),
    size=28, bold=True, color=NAVY)

txt(sl, 'Every health record that existed before this one was owned by a clinic, a hospital, or a health '
        'system \u2014 and could only be read, written, or moved with that institution\u2019s permission. '
        'My Shared Care Plan is different at its foundation: a federated, decentralized record held by '
        'the person, not by any institution.',
    Inches(0.7), Inches(1.65), Inches(11.5), Inches(1.0),
    size=14, color=TEXT_MUTED)

# Two cards
CW = Inches(5.7)
cy = Inches(2.85)

card(sl, Inches(0.7), cy, CW, Inches(2.1), TEAL_LT, TEAL_BD)
txt(sl, 'INTERFACED \u2014 NOT OWNED', Inches(0.86), cy + Inches(0.12),
    CW - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'Connects to the medical system when needed',
    Inches(0.86), cy + Inches(0.38), CW - Inches(0.32), Inches(0.38),
    size=13, bold=True, color=NAVY)
txt(sl, 'When a problem requires physician attention, the record speaks clearly to that. '
        'Every piece of evidence, every gap, the specific question being asked. The medical '
        'system receives exactly what it needs \u2014 and nothing more.',
    Inches(0.86), cy + Inches(0.82), CW - Inches(0.32), Inches(1.1),
    size=11, color=TEXT_MED)

card(sl, Inches(6.93), cy, CW, Inches(2.1), BLUE_LT, BLUE_BD)
txt(sl, 'PERSON-OWNED \u2014 ALWAYS', Inches(7.09), cy + Inches(0.12),
    CW - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'Belongs to the person, period',
    Inches(7.09), cy + Inches(0.38), CW - Inches(0.32), Inches(0.38),
    size=13, bold=True, color=NAVY)
txt(sl, 'No institution holds it. No clinic owns it. The person carries it into every '
        'relationship \u2014 with their health partner, their physician, their care team, '
        'their family. It follows them, not the other way around.',
    Inches(7.09), cy + Inches(0.82), CW - Inches(0.32), Inches(1.1),
    size=11, color=TEXT_MED)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 5 — Larry's Vision Realized
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, SURFACE)

label_tag(sl, "Weed\u2019s Problem Knowledge Coupler \u2014 realized", Inches(0.7))

txt(sl, 'Larry was pointing here.\nWe found the path he couldn\u2019t.',
    Inches(0.7), Inches(0.5), Inches(11.5), Inches(1.0),
    size=28, bold=True, color=NAVY)

txt(sl, 'Larry\u2019s vision was always that the person should own their health knowledge \u2014 structured, '
        'evidence-linked, actionable, accessible to anyone who was trusted to help. My Health Picture '
        'realizes his structured problem list. My Health Choices is similar in intent to the coupler \u2014 '
        'connecting each active problem to the relevant medical knowledge, options, and evidence. Together '
        'they are the direct realization of what Larry pointed toward, with everything that blocked the '
        'original designed out.',
    Inches(0.7), Inches(1.65), Inches(11.5), Inches(1.2),
    size=13, color=TEXT_MED)

# Three cards
CW3 = Inches(3.97)
cy = Inches(3.05)

card(sl, Inches(0.7), cy, CW3, Inches(2.2), SURFACE, BORDER)
txt(sl, 'WHAT LARRY HAD RIGHT', Inches(0.86), cy + Inches(0.12),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'The knowledge belongs at the point of care',
    Inches(0.86), cy + Inches(0.38), CW3 - Inches(0.32), Inches(0.42),
    size=12, bold=True, color=NAVY)
txt(sl, 'Problems stated at their level of certainty. Evidence linked to each diagnosis. '
        'Gaps made explicit. The person and their helper as the center of the record.',
    Inches(0.86), cy + Inches(0.86), CW3 - Inches(0.32), Inches(1.2),
    size=10, color=TEXT_MED)

card(sl, Inches(4.68), cy, CW3, Inches(2.2), TEAL_LT, TEAL_BD)
txt(sl, 'WHAT WE ADDED', Inches(4.84), cy + Inches(0.12),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'Open source \u00b7 person-owned \u00b7 CHW-operable',
    Inches(4.84), cy + Inches(0.38), CW3 - Inches(0.32), Inches(0.42),
    size=12, bold=True, color=NAVY)
txt(sl, 'No IP. No license. No institutional owner. A community health worker can operate '
        'these tools in a person\u2019s home today \u2014 no physician required to run them.',
    Inches(4.84), cy + Inches(0.86), CW3 - Inches(0.32), Inches(1.2),
    size=10, color=TEXT_MED)

card(sl, Inches(8.66), cy, CW3, Inches(2.2), GREEN_LT, GREEN_BD)
txt(sl, 'THE MISSING PIECE LARRY NEVER HAD', Inches(8.82), cy + Inches(0.12),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'A cooperative that pays for health',
    Inches(8.82), cy + Inches(0.38), CW3 - Inches(0.32), Inches(0.42),
    size=12, bold=True, color=NAVY)
txt(sl, 'When the tools live inside a system that earns more as problems are resolved, '
        'using them is how you get paid better. The incentive and the tools finally point '
        'the same direction.',
    Inches(8.82), cy + Inches(0.86), CW3 - Inches(0.32), Inches(1.2),
    size=10, color=TEXT_MED)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 6 — Section: The Tools
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, DARK_TEAL)

txt(sl, 'THE TOOLS',
    Inches(0.7), Inches(2.5), Inches(11.93), Inches(0.4),
    size=12, bold=True, color=RGBColor(0x80, 0xb8, 0xc8),
    align=PP_ALIGN.CENTER)

txt(sl, 'My Personal Health Supporter',
    Inches(0.7), Inches(3.0), Inches(11.93), Inches(1.0),
    size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

txt(sl, 'Everything a person needs  \u00b7  In their hands  \u00b7  At home',
    Inches(0.7), Inches(4.1), Inches(11.93), Inches(0.4),
    size=17, color=RGBColor(0x80, 0xb8, 0xc8), align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 7 — Three Tools
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, SURFACE)

label_tag(sl, 'My Personal Health Supporter', Inches(0.7))

txt(sl, 'Three tools. One complete\nperson-centered picture.',
    Inches(0.7), Inches(0.48), Inches(11.5), Inches(1.0),
    size=28, bold=True, color=NAVY)

CW3 = Inches(3.97)
cy = Inches(1.65)

# Card 1: My Shared Care Plan (purple)
card(sl, Inches(0.7), cy, CW3, Inches(4.5), PURPLE_LT, PURPLE_BD)
txt(sl, 'MY SHARED CARE PLAN', Inches(0.86), cy + Inches(0.15),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'The medical record',
    Inches(0.86), cy + Inches(0.42), CW3 - Inches(0.32), Inches(0.35),
    size=14, bold=True, color=NAVY)
txt(sl, 'The actual clinical record \u2014 diagnoses, medications, vitals, care team, visits, '
        'directives, history, lab results. Clinical ground truth. Person-owned, decentralized, '
        'no institutional owner. Every other tool draws from this.',
    Inches(0.86), cy + Inches(0.85), CW3 - Inches(0.32), Inches(3.5),
    size=12, color=TEXT_MED)

# Card 2: My Health Picture (blue)
card(sl, Inches(4.68), cy, CW3, Inches(4.5), BLUE_LT, BLUE_BD)
txt(sl, 'MY HEALTH PICTURE', Inches(4.84), cy + Inches(0.15),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'The active problem view',
    Inches(4.84), cy + Inches(0.42), CW3 - Inches(0.32), Inches(0.35),
    size=14, bold=True, color=NAVY)
txt(sl, 'Organized by Subjective, Objective, Assessment, and Plan. Draws from the SCP to show '
        'the active problem list with three flags per problem: degree of certainty, urgency, and '
        'missing data. What the health partner and person review together before every encounter.',
    Inches(4.84), cy + Inches(0.85), CW3 - Inches(0.32), Inches(3.5),
    size=12, color=TEXT_MED)

# Card 3: My Health Choices (teal)
card(sl, Inches(8.66), cy, CW3, Inches(4.5), TEAL_LT, TEAL_BD)
txt(sl, 'MY HEALTH CHOICES', Inches(8.82), cy + Inches(0.15),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'The option explorer',
    Inches(8.82), cy + Inches(0.42), CW3 - Inches(0.32), Inches(0.35),
    size=14, bold=True, color=NAVY)
txt(sl, 'AI-assisted exploration of a specific problem against the medical literature. What are '
        'the options? What does the evidence say? What are the tradeoffs? Informed decision-making '
        'at home, before the appointment \u2014 not a signature after the fact.',
    Inches(8.82), cy + Inches(0.85), CW3 - Inches(0.32), Inches(3.5),
    size=12, color=TEXT_MED)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 8 — The Health Partner
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, GREY_BG)

label_tag(sl, 'The neighborhood figure', Inches(0.7))

txt(sl, 'A health partner \u2014\nfrom the community, in the home',
    Inches(0.7), Inches(0.48), Inches(11.5), Inches(1.0),
    size=28, bold=True, color=NAVY)

txt(sl, 'The health partner is not a healthcare extender. They are a neighborhood figure with real tools \u2014 '
        'someone who knows the person, knows the street, and carries the full structured record into every '
        'encounter. The Southcentral Foundation proved what this looks like at scale: community health aides '
        'at the center of an Alaska Native-owned system, achieving results no clinic-centered model has matched.',
    Inches(0.7), Inches(1.65), Inches(11.5), Inches(1.1),
    size=14, color=TEXT_MUTED)

# Flow row
flow_y = Inches(3.0)
FW = Inches(2.5)
fa_x = [Inches(0.7), Inches(3.38), Inches(6.06), Inches(8.74)]
fa_labels = [
    ('Health partner', 'uses the tools at home'),
    ('Structured package', 'problem + evidence + question'),
    ('Physician response', 'to a precise question'),
    ('Shared Care Plan', 'written, visible to all'),
]

for i, (main, sub) in enumerate(fa_labels):
    bx = fa_x[i]
    fill = TEAL_LT if i < 2 else GREEN_LT
    bd   = TEAL_BD if i < 2 else GREEN_BD
    card(sl, bx, flow_y, FW, Inches(0.85), fill, bd)
    txt(sl, main, bx + Inches(0.1), flow_y + Inches(0.08),
        FW - Inches(0.2), Inches(0.32), size=11, bold=True, color=NAVY)
    txt(sl, sub, bx + Inches(0.1), flow_y + Inches(0.4),
        FW - Inches(0.2), Inches(0.3), size=9, color=TEXT_MUTED)
    if i < 3:
        txt(sl, '\u2192', bx + FW + Inches(0.06), flow_y + Inches(0.26),
            Inches(0.28), Inches(0.32), size=16, bold=True, color=TEAL)

txt(sl, 'Southcentral Foundation  \u00b7  Nuka System of Care  \u00b7  \u221236% ED visits  \u00b7  '
        '\u221240% specialist visits  \u00b7  sustained over decades',
    Inches(0.7), Inches(4.1), Inches(11.93), Inches(0.35),
    size=11, color=TEXT_GREY, align=PP_ALIGN.CENTER)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 9 — RenDanHeYi
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, NAVY)

txt(sl, 'The organizing philosophy  \u00b7  \u4eba\u5355\u5408\u4e00',
    Inches(0.7), Inches(0.45), Inches(11.93), Inches(0.45),
    size=18, bold=True, color=RGBColor(0x88, 0x99, 0xaa),
    align=PP_ALIGN.CENTER)

# Four cards
CW4 = Inches(2.96)
cy = Inches(1.1)
chars = [('\u4eba', 'R\u00e9n', 'Person. Not a role, not a unit of capacity \u2014 a moral agent with dignity and initiative. The patient, the health partner, the physician: each a full person.'),
         ('\u5355', 'D\u0101n', 'Order. A specific, concrete need \u2014 not \u201cchronic disease management\u201d but this person\u2019s open problem, today, in their home.'),
         ('\u5408', 'H\u00e9', 'Active merging. The living process of person and need coming together \u2014 not yet resolved, but in genuine motion. The care relationship itself.'),
         ('\u4e00', 'Y\u012b', 'Achieved oneness. Resolution. The problem actually addressed. What the cooperative is paid for \u2014 not the encounter, but the \u4e00.')]

for i, (hanzi, roman, meaning) in enumerate(chars):
    cx = Inches(0.7) + i * (CW4 + Inches(0.2))
    shp = sl.shapes.add_shape(1, cx, cy, CW4, Inches(3.6))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(0x28, 0x42, 0x62)
    shp.line.color.rgb = RGBColor(0x44, 0x5a, 0x7a)
    shp.line.width = Pt(1)

    txt(sl, hanzi, cx, cy + Inches(0.15), CW4, Inches(0.9),
        size=44, color=WHITE, align=PP_ALIGN.CENTER)
    txt(sl, roman, cx, cy + Inches(1.05), CW4, Inches(0.3),
        size=11, color=RGBColor(0x88, 0x99, 0xcc), align=PP_ALIGN.CENTER)
    txt(sl, meaning, cx + Inches(0.15), cy + Inches(1.4), CW4 - Inches(0.3), Inches(2.0),
        size=11, color=RGBColor(0xcc, 0xd5, 0xe0))

# Synthesis box
sy = Inches(4.9)
shp2 = sl.shapes.add_shape(1, Inches(0.7), sy, Inches(11.93), Inches(1.2))
shp2.fill.solid()
shp2.fill.fore_color.rgb = RGBColor(0x24, 0x3a, 0x55)
shp2.line.color.rgb = RGBColor(0x3a, 0x55, 0x7a)
shp2.line.width = Pt(1)
txt(sl, 'The person and the order are in the active process of becoming one. The neighborhood '
        'microenterprise earns when \u4e00 is achieved \u2014 which means the health partner, the tools, '
        'and the economics all point toward the same thing: health.',
    Inches(0.9), sy + Inches(0.18), Inches(11.53), Inches(0.9),
    size=14, italic=True, color=RGBColor(0xcc, 0xd8, 0xee), align=PP_ALIGN.CENTER)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 10 — The Cooperative
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, SURFACE)

label_tag(sl, 'Whatcom Wealth and Health Association', Inches(0.7))

txt(sl, 'A neighborhood economic structure \u2014\norganized around what neighborhoods do',
    Inches(0.7), Inches(0.48), Inches(11.5), Inches(1.0),
    size=28, bold=True, color=NAVY)

txt(sl, 'The cooperative is not a healthcare improvement project. It is an economic structure that makes '
        'the neighborhood health system financially sustainable and self-reinforcing. Large Whatcom employers '
        'and governments enter as self-insured members. Health partner microenterprises \u2014 self-managing '
        'teams of 5 to 15 people from the neighborhoods they serve \u2014 are paid for problems resolved and '
        'health maintained.',
    Inches(0.7), Inches(1.65), Inches(11.5), Inches(1.2),
    size=13, color=TEXT_MED)

CW3 = Inches(3.97)
cy = Inches(3.05)

card(sl, Inches(0.7), cy, CW3, Inches(2.5), TEAL_LT, TEAL_BD)
txt(sl, 'ENTRY POINT', Inches(0.86), cy + Inches(0.12),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'Employer self-insurance',
    Inches(0.86), cy + Inches(0.38), CW3 - Inches(0.32), Inches(0.35),
    size=13, bold=True, color=NAVY)
txt(sl, 'Businesses and governments who self-insure through the cooperative gain fiduciary benefits '
        'design, direct primary care, and transparent pricing. Dave Chase\u2019s Health Rosetta model '
        '\u2014 20\u201340% cost reduction documented.',
    Inches(0.86), cy + Inches(0.82), CW3 - Inches(0.32), Inches(1.55),
    size=11, color=TEXT_MED)

card(sl, Inches(4.68), cy, CW3, Inches(2.5), GREEN_LT, GREEN_BD)
txt(sl, 'THE CARE UNIT', Inches(4.84), cy + Inches(0.12),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'Neighborhood microenterprise',
    Inches(4.84), cy + Inches(0.38), CW3 - Inches(0.32), Inches(0.35),
    size=13, bold=True, color=NAVY)
txt(sl, '5\u201315 health partners, self-managing, their own P&L, paid for outcomes. Physician '
        'partners contracted for specific expertise. No management above them \u2014 only the '
        'commitments they have made.',
    Inches(4.84), cy + Inches(0.82), CW3 - Inches(0.32), Inches(1.55),
    size=11, color=TEXT_MED)

card(sl, Inches(8.66), cy, CW3, Inches(2.5), BLUE_LT, BLUE_BD)
txt(sl, 'OPEN TO ALL', Inches(8.82), cy + Inches(0.12),
    CW3 - Inches(0.32), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'The whole population',
    Inches(8.82), cy + Inches(0.38), CW3 - Inches(0.32), Inches(0.35),
    size=13, bold=True, color=NAVY)
txt(sl, 'Starts with employers. Extends to every person in Whatcom County who chooses to join. '
        'Any member can become a health partner. The cooperative grows stronger as it grows larger.',
    Inches(8.82), cy + Inches(0.82), CW3 - Inches(0.32), Inches(1.55),
    size=11, color=TEXT_MED)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 11 — Open Source
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, GREY_BG)

label_tag(sl, 'What happened to Larry\u2019s tools \u2014 and what we did about it', Inches(0.7))

txt(sl, 'Critical health infrastructure\nmust belong to the commons',
    Inches(0.7), Inches(0.48), Inches(11.5), Inches(1.0),
    size=28, bold=True, color=NAVY)

txt(sl, 'When Larry\u2019s IP was sold, distribution was strangled and the tools died with the ownership '
        'complications. That is the cautionary case. Every tool I build is and will continue to be open '
        'source, Creative Commons Attribution 4.0 \u2014 permanently, structurally, by design. This is '
        'not a policy choice. It is the founding architecture.',
    Inches(0.7), Inches(1.65), Inches(11.5), Inches(1.1),
    size=14, color=TEXT_MUTED)

# Open source box
bx = Inches(0.7)
by = Inches(3.0)
bw = Inches(11.93)
bh = Inches(2.4)
shp = sl.shapes.add_shape(1, bx, by, bw, bh)
shp.fill.solid()
shp.fill.fore_color.rgb = GREEN_BOX
shp.line.color.rgb = GREEN_BR
shp.line.width = Pt(2)

txt(sl, 'PERMANENT COMMITMENT',
    bx, by + Inches(0.2), bw, Inches(0.28),
    size=10, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)

txt(sl, 'Every RCN tool is open source CC 4.0. None are for sale. None ever will be.\n'
        'The cooperative is the financial sustainability model. The tools belong to the commons.',
    bx + Inches(0.5), by + Inches(0.55), bw - Inches(1.0), Inches(0.75),
    size=16, bold=True, color=GREEN_TEXT, align=PP_ALIGN.CENTER)

txt(sl, 'Open infrastructure + cooperative economics = a system that cannot be captured by any single owner.\n'
        'The person\u2019s record, the neighborhood\u2019s tools, the community\u2019s health \u2014 held in common.',
    bx + Inches(0.5), by + Inches(1.4), bw - Inches(1.0), Inches(0.8),
    size=13, color=GREEN_DARK, align=PP_ALIGN.CENTER)

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 12 — Where Things Stand
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, NAVY)

label_txt = 'A DISPATCH FROM THE WORK'
txt(sl, label_txt, Inches(0.7), Inches(0.42), Inches(11.5), Inches(0.28),
    size=9, bold=True, color=RGBColor(0x77, 0x88, 0xaa))

txt(sl, 'Where things stand \u2014\nJuly 2026',
    Inches(0.7), Inches(0.72), Inches(11.5), Inches(0.95),
    size=28, bold=True, color=WHITE)

CW2 = Inches(5.7)
cy = Inches(1.85)

# Card: The Tools (navy on navy)
shp1 = sl.shapes.add_shape(1, Inches(0.7), cy, CW2, Inches(2.3))
shp1.fill.solid()
shp1.fill.fore_color.rgb = NAVY_LT
shp1.line.color.rgb = RGBColor(0x3a, 0x55, 0x7a)
shp1.line.width = Pt(1)
txt(sl, 'THE TOOLS', Inches(0.86), cy + Inches(0.12),
    CW2 - Inches(0.32), Inches(0.22), size=9, bold=True, color=RGBColor(0x77, 0x99, 0xbb))
txt(sl, 'Built and working',
    Inches(0.86), cy + Inches(0.38), CW2 - Inches(0.32), Inches(0.35),
    size=13, bold=True, color=WHITE)
txt(sl, 'My Health Picture, My Shared Care Plan, My Health Choices, My Support Network, and SODOTO are '
        'open source, running, and in use. The person\u2019s record is untethered. Appointment prep, '
        'PANAS activation scoring, and AI-assisted option exploration are active today.',
    Inches(0.86), cy + Inches(0.82), CW2 - Inches(0.32), Inches(1.35),
    size=11, color=RGBColor(0xb0, 0xc8, 0xe8))

# Card: The Cooperative
shp2 = sl.shapes.add_shape(1, Inches(6.93), cy, CW2, Inches(2.3))
shp2.fill.solid()
shp2.fill.fore_color.rgb = NAVY_LT
shp2.line.color.rgb = RGBColor(0x3a, 0x55, 0x7a)
shp2.line.width = Pt(1)
txt(sl, 'THE COOPERATIVE', Inches(7.09), cy + Inches(0.12),
    CW2 - Inches(0.32), Inches(0.22), size=9, bold=True, color=RGBColor(0x77, 0x99, 0xbb))
txt(sl, 'Organizing now',
    Inches(7.09), cy + Inches(0.38), CW2 - Inches(0.32), Inches(0.35),
    size=13, bold=True, color=WHITE)
txt(sl, 'Whatcom Wealth and Health Association \u2014 Marc, Dave Chase, a local insurance broker, '
        'two Whatcom community business leaders \u2014 actively organizing employer self-insurance '
        'entry with large local employers and county government.',
    Inches(7.09), cy + Inches(0.82), CW2 - Inches(0.32), Inches(1.35),
    size=11, color=RGBColor(0xb0, 0xc8, 0xe8))

# Dispatch box
dy = Inches(4.35)
shp3 = sl.shapes.add_shape(1, Inches(0.7), dy, Inches(11.93), Inches(1.45))
shp3.fill.solid()
shp3.fill.fore_color.rgb = RGBColor(0xeb, 0xfd, 0xfb)
shp3.line.color.rgb = DISPATCH_BD
shp3.line.width = Pt(2)
txt(sl, 'THE PATH IS CLEAR', Inches(0.86), dy + Inches(0.1),
    Inches(11.61), Inches(0.22), size=9, bold=True, color=TEAL)
txt(sl, 'There have been unexpected obstacles, as you know. But the design is right, the tools are real, '
        'the cooperative structure sidesteps everything that stopped WIDS in 1995. We are not asking the '
        'system to change. We are building alongside it. Others will follow when they see what is possible.',
    Inches(0.86), dy + Inches(0.38), Inches(11.61), Inches(0.95),
    size=13, color=RGBColor(0x0c, 0x44, 0x44))

footer(sl)

# ═══════════════════════════════════════════════════════════════════════════════
# Slide 13 — Closing Quote
# ═══════════════════════════════════════════════════════════════════════════════
sl = slide()
bg(sl, NAVY)

txt(sl, '\u201cGive people and their communities everything they need and want \u2014\n'
        'the knowledge, the tools, the economic support \u2014\n'
        'and watch what human beings do with their own health.\u201d',
    Inches(0.9), Inches(1.8), Inches(11.53), Inches(2.4),
    size=26, italic=True, color=WHITE, align=PP_ALIGN.CENTER)

txt(sl, 'Whatcom County  \u00b7  2026  \u00b7  Open Source CC 4.0  \u00b7  Marc Pierson MD',
    Inches(0.7), Inches(4.5), Inches(11.93), Inches(0.35),
    size=12, color=RGBColor(0x77, 0x88, 0xaa), align=PP_ALIGN.CENTER)

# ── Save ──────────────────────────────────────────────────────────────────────
prs.save(OUT)
print(f'Saved: {OUT}')
print(f'Slides: {len(prs.slides)}')
