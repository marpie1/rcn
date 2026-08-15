"""
Generate tools/rcn-bias-checker-intro.pptx — the Conversation Navigator intro deck.

Companion to tools/bias-checker-intro.html and bias-checker-manual.html; same
argument, deck shape. Regenerate after editing:
    python3 docs/make_bias_checker_pptx.py
Requires: pip install python-pptx

The tool's filename is bias-checker.html for historical reasons; its name is the
Conversation Navigator. The deck uses the name, the repo uses the filename.

Content here is the real vocabulary of the tool — the eight intents, the six
biases, the conflict table — copied from the INTENTS / BIASES / CONFLICTS
constants in bias-checker.html. If those change, change these to match.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette — the tool's own warm ink-and-terracotta, not the generic doc blue ─
ACC    = RGBColor(0xc8, 0x4b, 0x2f)   # terracotta — this tool
ACC_LT = RGBColor(0xfb, 0xe4, 0xde)
ACC_DK = RGBColor(0x8a, 0x2a, 0x14)
AMB    = RGBColor(0xe8, 0xa0, 0x20)
AMB_LT = RGBColor(0xfd, 0xf3, 0xdc)
GRN    = RGBColor(0x3a, 0x7d, 0x44)
GRN_LT = RGBColor(0xe0, 0xef, 0xe2)
INK    = RGBColor(0x1a, 0x18, 0x14)
INK2   = RGBColor(0x5a, 0x56, 0x50)
INK3   = RGBColor(0x9a, 0x96, 0x90)
SURF2  = RGBColor(0xed, 0xea, 0xe2)
BORDER = RGBColor(0xd8, 0xd4, 0xcc)
BG     = RGBColor(0xf5, 0xf2, 0xeb)
WHITE  = RGBColor(0xff, 0xff, 0xff)

INTENTS = [
    ("Learning",          "Understand something you do not already know"),
    ("Socializing",       "The relationship is the point"),
    ("Convincing",        "You hold a position and want it adopted"),
    ("Exploring",         "The question is genuinely open"),
    ("Producing",         "There is a deliverable"),
    ("Venting",           "Discharging energy; content is secondary"),
    ("Being heard",       "You need to feel understood first"),
    ("Sharing / Teaching","You have something to transfer"),
]

BIASES = [
    ("Divided attention",          "Part of you is elsewhere"),
    ("Assuming success",           "You believe you were understood — unverified"),
    ("Doubling down",              "Pushback hardens rather than opens you"),
    ("Statements : Questions > 1", "Asserting more than asking"),
    ("Tone limiting participation","Your manner makes it harder for others"),
    ("We need a model",            "Reaching for a framework instead of specifics"),
]

CONFLICTS = [
    ("Learning",          "Assuming success · Doubling down"),
    ("Socializing",       "Doubling down · Tone limiting"),
    ("Convincing",        "Divided attention · Assuming success"),
    ("Exploring",         "Doubling down · We need a model"),
    ("Producing",         "Divided attention · Doubling down · We need a model"),
    ("Venting",           "(none)"),
    ("Being heard",       "Doubling down · Tone limiting"),
    ("Sharing / Teaching","Statements > Questions · Tone limiting"),
]

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Primitives ───────────────────────────────────────────────────────────────

def slide():
    return prs.slides.add_slide(BLANK)


def bg(sl, color):
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = color


def rect(sl, x, y, w, h, fill=WHITE, line=None, lw=1.5):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def txt(sl, text, x, y, w, h, size=20, bold=False, italic=False,
        color=INK, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return tb


def txlines(sl, lines, x, y, w, h, sizes=None, bolds=None, colors=None,
            align=PP_ALIGN.LEFT, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(sizes[i] if isinstance(sizes, list) else (sizes or 18))
        r.font.bold = bolds[i] if isinstance(bolds, list) else (bolds or False)
        r.font.color.rgb = colors[i] if isinstance(colors, list) else (colors or INK)
        r.font.name = 'Calibri'
    return tb


def box_txt(sl, lines, x, y, w, h, fill=BG, line=BORDER, lw=1.5,
            size=14, bold_first=True, text_color=INK2,
            top_color=INK, align=PP_ALIGN.LEFT):
    sh = rect(sl, x, y, w, h, fill=fill, line=line, lw=lw)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.11)
    tf.margin_bottom = Inches(0.1)
    for i, line_text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(size if i > 0 else size + 1)
        r.font.bold = (bold_first and i == 0)
        r.font.color.rgb = (top_color or text_color) if i == 0 else text_color
        r.font.name = 'Calibri'
    return sh


def label(sl, text, color=ACC_DK):
    txt(sl, text, 0.45, 0.18, 8, 0.3, size=9, bold=True, color=color)


def bar(sl, color=ACC, y=0.55, h=0.035):
    rect(sl, 0, y, 13.33, h, fill=color)


def title_txt(sl, text, color=INK, size=34, y=0.62):
    txt(sl, text, 0.45, y, 12.4, 1.3, size=size, bold=True, color=color)


def foot(sl, text):
    txt(sl, text, 0.45, 6.95, 12.4, 0.35, size=10, color=INK3)


def head(text, kicker="CONVERSATION NAVIGATOR"):
    s = slide()
    bg(s, WHITE)
    label(s, kicker)
    bar(s)
    title_txt(s, text)
    return s


# ── 1 · Title ────────────────────────────────────────────────────────────────
s = slide()
bg(s, INK)
rect(s, 0, 0, 13.33, 0.16, fill=ACC)
txt(s, "RELOCALIZE CREATIVITY NETWORK · OPEN SOURCE · FREE BY DESIGN",
    0.9, 1.35, 11.5, 0.4, size=12, bold=True, color=INK3)
txt(s, "Conversation Navigator", 0.9, 1.95, 11.5, 1.2, size=54, bold=True, color=WHITE)
txlines(s, [
    "Make the intent and the biases of everyone in the room",
    "visible to each other — during the conversation,",
    "not in the debrief afterwards.",
], 0.9, 3.35, 11.5, 2.0, sizes=24, colors=ACC_LT, spacing=6)
txt(s, "Runs in a browser · no account · one shared link", 0.9, 5.5, 11.5, 0.5,
    size=17, italic=True, color=INK3)
txt(s, "tools/bias-checker.html", 0.9, 6.3, 11, 0.4, size=15, color=ACC_LT)

# ── 2 · Why it exists ────────────────────────────────────────────────────────
s = head("Most facilitation breakdowns are not failures of method")
txt(s, "They are failures of transparency.", 0.45, 1.85, 12.4, 0.5,
    size=22, italic=True, color=ACC_DK)
box_txt(s, [
    "Everyone arrives with an intent",
    "One person came to decide. Another came to think out loud. A third came to be heard. "
    "Nobody says so, and the meeting is judged a failure by three different standards.",
], 0.45, 2.6, 6.1, 1.75, fill=BG, size=15)
box_txt(s, [
    "Everyone arrives with a bias",
    "Half-attending. Certain of being understood. Hardening under challenge. These are ordinary, "
    "and mostly invisible to the person having them.",
], 6.78, 2.6, 6.1, 1.75, fill=BG, size=15)
box_txt(s, [
    "The tool does one thing",
    "It puts both on the table, live, where the group can see them — and flags when a person's own "
    "stated intent is in tension with a bias they have just admitted to.",
], 0.45, 4.6, 12.43, 1.5, fill=ACC_LT, line=ACC, size=15, top_color=ACC_DK)
foot(s, "This is a self-report tool. You declare your own intent and notice your own biases — never anyone else's.")

# ── 3 · How a session works ──────────────────────────────────────────────────
s = head("Four steps, no account")
steps = [
    ("1 · Start",  "Open the tool. A six-character session ID is generated and added to the URL."),
    ("2 · Invite", "Share that URL. Anyone who opens it is in the same live session."),
    ("3 · Join",   "Enter a name. Your card appears for everyone, and updates as you change it."),
    ("4 · Close",  "Export the session, then end it. Ending deletes everything."),
]
for i, (t, d) in enumerate(steps):
    box_txt(s, [t, d], 0.45 + i * 3.19, 2.0, 2.98, 1.9, fill=BG, size=14)
box_txt(s, [
    "Update it during the conversation, not just at the start",
    "The point is not the opening declaration. It is noticing, twenty minutes in, that you came to "
    "explore and have started convincing — and changing your card while the others watch.",
], 0.45, 4.3, 12.43, 1.5, fill=AMB_LT, line=AMB, size=15, top_color=RGBColor(0x92, 0x40, 0x0e))
foot(s, "Everything syncs live through a Firebase Realtime Database. No logins, no installs, no server of your own.")

# ── 4 · The eight intents ────────────────────────────────────────────────────
s = head("Your intent — pick as many as are true")
for i, (t, d) in enumerate(INTENTS):
    col, row = i % 4, i // 4
    box_txt(s, [t, d], 0.45 + col * 3.19, 1.95 + row * 1.55, 2.98, 1.35,
            fill=BG, size=13)
txt(s, "Add your own if none of these fit. Drag to reorder. Multiple intents are normal — "
       "\"being heard\" and \"producing something\" are often both live at once.",
    0.45, 5.25, 12.43, 0.9, size=15, color=INK2)
foot(s, "Eight built-in intents; custom intents are per-participant and appear in the group stats alongside the built-ins.")

# ── 5 · The six biases ───────────────────────────────────────────────────────
s = head("Biases I am noticing in myself")
for i, (t, d) in enumerate(BIASES):
    col, row = i % 3, i // 3
    box_txt(s, [t, d], 0.45 + col * 4.25, 1.95 + row * 1.5, 4.04, 1.3,
            fill=BG, size=13)
box_txt(s, [
    "The wording is doing real work",
    "Not \"biases in the room\" and not \"biases I see in you\". A tool that let people tag each other "
    "would be a weapon. This one only lets you tag yourself, which is the only version a group will "
    "actually use twice.",
], 0.45, 5.05, 12.43, 1.5, fill=ACC_LT, line=ACC, size=15, top_color=ACC_DK)
foot(s, "Six built-in biases; you can add your own.")

# ── 6 · Conflict detection ───────────────────────────────────────────────────
s = head("Where the tool speaks up")
txt(s, "When your stated intent conflicts with a bias you have just admitted, a warning appears on "
       "your card — visible to everyone.", 0.45, 1.8, 12.4, 0.6, size=17, color=INK2)
rect(s, 0.45, 2.55, 12.43, 0.42, fill=INK)
txt(s, "INTENT", 0.62, 2.6, 3, 0.32, size=11, bold=True, color=BG)
txt(s, "CONFLICTS WITH", 4.3, 2.6, 6, 0.32, size=11, bold=True, color=BG)
for i, (intent, clash) in enumerate(CONFLICTS):
    y = 2.97 + i * 0.44
    rect(s, 0.45, y, 12.43, 0.44, fill=(WHITE if i % 2 == 0 else BG), line=BORDER, lw=0.5)
    txt(s, intent, 0.62, y + 0.06, 3.6, 0.32, size=13, bold=True, color=INK)
    txt(s, clash, 4.3, y + 0.06, 8.3, 0.32, size=13,
        color=(INK3 if clash == "(none)" else ACC_DK))
foot(s, "The table is deterministic — the tool never infers a conflict, it looks one up.")

# ── 7 · Positive Affect Score ────────────────────────────────────────────────
s = head("Positive Affect Score — how much capacity is actually in the room")
txt(s, "Ten words, rated 1–5, adding to a score between 10 and 50. Under thirty seconds.",
    0.45, 1.8, 12.4, 0.5, size=18, color=INK2)
box_txt(s, [
    "Take it twice",
    "Once at the start, once at the end. Your card shows your own shift; the Stats view shows how the "
    "whole room moved. The change is the number worth looking at — a single score in isolation says "
    "very little.",
], 0.45, 2.5, 6.1, 2.1, fill=ACC_LT, line=ACC, size=15, top_color=ACC_DK)
box_txt(s, [
    "Deliberately no bands here",
    "The four activation bands used in the health tools are calibrated for a past-week reading. "
    "Labelling someone in a meeting \"Learning the basics\" on the strength of their last twenty "
    "minutes would be both wrong and insulting.",
], 6.78, 2.5, 6.1, 2.1, fill=BG, size=15)
box_txt(s, [
    "Optional NA-10",
    "A second set of ten negative-affect words, scored separately and never subtracted from the "
    "positive score. The two are close to independent — a person can be high on both at once.",
], 0.45, 4.82, 12.43, 1.4, fill=BG, size=15)
foot(s, "PANAS: Watson, Clark & Tellegen (1988). The standalone version with history, past-week framing and bands is tools/positive-affect-score.html")

# ── 8 · Shared models ────────────────────────────────────────────────────────
s = head("Five shared models everyone can click")
models = [
    ("Six Questions", "Edit the questions, type answers, vote on each other's"),
    ("Cynefin",       "Mark the domain, or the boundary you are crossing"),
    ("eVSM v2",       "Click the function or the flow under discussion"),
    ("15 Ps",         "Mark what is relevant to the topic"),
    ("Six Hats",      "Which mode is actually dominant right now"),
]
for i, (t, d) in enumerate(models):
    box_txt(s, [t, d], 0.45 + i * 2.53, 2.0, 2.36, 1.7, fill=BG, size=13)
box_txt(s, [
    "Counts aggregate live, and hovering shows who",
    "Every click is attributed. In an open session, hovering any bar in the Stats view names the "
    "people behind it — intents, biases, model clicks, affect scores and Six Questions votes alike. "
    "Upload a graph JSON from the RCN Graph Tool to add a diagram of your own.",
], 0.45, 4.05, 12.43, 1.7, fill=ACC_LT, line=ACC, size=15, top_color=ACC_DK)
foot(s, "Models are shared click targets — unlike intents, biases and the affect score, which are private self-reports about yourself.")

# ── 9 · Identity and consent ─────────────────────────────────────────────────
s = head("Names are on by default — and anyone can switch them off")
box_txt(s, [
    "Open",
    "Names are visible to everyone. Hover any aggregate to see who holds that intent, "
    "checked that bias, clicked that node or scored that number.",
], 0.45, 2.0, 6.1, 1.7, fill=GRN_LT, line=GRN, size=15, top_color=RGBColor(0x2a, 0x5c, 0x33))
box_txt(s, [
    "Anonymous",
    "The moment one participant ticks the box, the whole session goes anonymous for everyone. "
    "Distributions, means and counts remain; no name appears anywhere.",
], 6.78, 2.0, 6.1, 1.7, fill=SURF2, line=INK2, size=15)
box_txt(s, [
    "One person is enough, and that is the design",
    "Anonymity is not a majority vote. If a single person in the room needs cover to answer "
    "honestly, they get it without having to argue for it — and without having to identify "
    "themselves as the one who asked.",
], 0.45, 4.0, 12.43, 1.6, fill=ACC_LT, line=ACC, size=15, top_color=ACC_DK)
foot(s, "Someone who leaves keeps their anonymity choice: their data is still in the export, so their consent still binds.")

# ── 10 · Leaving vs removing ─────────────────────────────────────────────────
s = head("Leaving does not withdraw what you said")
box_txt(s, [
    "Close your tab",
    "Your card greys out and is marked LEFT. Everything you contributed stays in the session and in "
    "its export. A crash or a sleeping phone behaves the same way.",
], 0.45, 2.0, 4.04, 1.9, fill=BG, size=14)
box_txt(s, [
    "Remove my input",
    "Deletes your card and your contributions for good. This is the only thing that erases them, and "
    "it asks first.",
], 4.65, 2.0, 4.04, 1.9, fill=AMB_LT, line=AMB, size=14,
    top_color=RGBColor(0x92, 0x40, 0x0e))
box_txt(s, [
    "End session",
    "Clears the whole session — every record, topic and model interaction. Export first; it cannot "
    "be undone.",
], 8.85, 2.0, 4.03, 1.9, fill=ACC_LT, line=ACC, size=14, top_color=ACC_DK)
box_txt(s, [
    "Why it works this way",
    "If leaving deleted your input, an export would only ever capture whoever happened to still have "
    "a tab open — which is to say, the least representative version of the meeting. Rejoining under "
    "the same name resumes your existing card rather than creating a second one.",
], 0.45, 4.2, 12.43, 1.6, fill=WHITE, line=BORDER, size=15)
foot(s, "A session outlives the meeting. Anyone with the six-character link can read it until somebody ends it — so end it when the group is finished.")

# ── 11 · Export ──────────────────────────────────────────────────────────────
s = head("What you take away")
box_txt(s, [
    "FedWiki page JSON",
    "Drags straight onto any FedWiki server. Participants, intent and bias distributions, affect "
    "scores, topics, model interactions, Six Questions answers — plus your own summary and "
    "carry-forward notes.",
], 0.45, 2.0, 6.1, 2.1, fill=BG, size=15)
box_txt(s, [
    "Raw session JSON",
    "The full state for archiving or your own analysis. Names are stripped automatically when the "
    "session is anonymous, exactly as they are everywhere else.",
], 6.78, 2.0, 6.1, 2.1, fill=BG, size=15)
box_txt(s, [
    "Two fields worth filling in before you export",
    "Summary — what happened, what shifted.    Carry Forward — what the group needs to bring next time.",
], 0.45, 4.35, 12.43, 1.3, fill=ACC_LT, line=ACC, size=15, top_color=ACC_DK)
foot(s, "Both exports include people who have already left.")

# ── 12 · Close ───────────────────────────────────────────────────────────────
s = slide()
bg(s, INK)
rect(s, 0, 0, 13.33, 0.14, fill=ACC)
txt(s, "THE BET UNDERNEATH", 0.9, 1.45, 11, 0.4, size=13, bold=True, color=ACC_LT)
txlines(s, [
    "A group cannot fix how it is talking",
    "while the way it is talking",
    "is the one thing nobody will say out loud.",
], 0.9, 2.1, 11.5, 2.4, sizes=30, bolds=True, colors=WHITE, spacing=8)
txt(s, "So make it the cheapest thing in the room to say.", 0.9, 4.8, 11.5, 0.5,
    size=19, italic=True, color=ACC_LT)
txt(s, "tools/bias-checker.html", 0.9, 5.75, 11, 0.5, size=18, color=ACC_LT)
txt(s, "Introduction · User Manual · Process Flow · hosted at Wiki Café",
    0.9, 6.25, 11.5, 0.4, size=13, color=INK3)

# ── Save ─────────────────────────────────────────────────────────────────────
here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, '..', 'tools', 'rcn-bias-checker-intro.pptx')
prs.save(out)
print('Wrote', os.path.normpath(out), '—', len(prs.slides._sldIdLst), 'slides')
