"""
Generate methods/kitchen-table/kitchen-table-exercise.pptx — the deck that runs
the Kitchen Table Exercise without a teacher in the room.

Companion to the FedWiki pages in methods/kitchen-table/ (hub, intro, eleven
step pages, notes sheet, manual). Regenerate after editing:
    python3 docs/make_kitchen_table_pptx.py
Requires: pip install python-pptx

AUDIENCE: two uses from one file. A person clicking through it alone, and a
host projecting it for a room where everyone works separately at the same
time. Both are served by the same slides because the host's only job is to
advance them.

DESIGN CONSTRAINT, deliberate and load bearing: slides 7-17 (the eleven steps)
carry the instruction and nothing else. No examples, no worked answer, no
partial diagram accumulating in a margin. A worked example is a hint in
written form -- it tells the person what a good answer looks like, which puts
a judge back in the room and converts the hour into a test. The house style
says every slide earns a visual; here the step slides earn theirs with a large
pale step numeral and white space, and all the design weight sits in the
framing slides before and the debrief slides after.

The fold slide is a hard stop. Everything after it assumes the work is done.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# -- Palette: graphite on paper. Chosen for a method whose only equipment is a
#    pencil. One accent (slate) for structure, one (oxide) for the two slides
#    that stop you.
INK    = RGBColor(0x14, 0x13, 0x0f)
INK2   = RGBColor(0x4a, 0x47, 0x40)
INK3   = RGBColor(0x91, 0x8c, 0x82)
PALE   = RGBColor(0xe3, 0xdf, 0xd6)
PAPER  = RGBColor(0xf6, 0xf4, 0xef)
RULE   = RGBColor(0xdd, 0xd8, 0xcd)
ACC    = RGBColor(0x3d, 0x5a, 0x6c)
ACC_LT = RGBColor(0xe4, 0xeb, 0xef)
OX     = RGBColor(0xa3, 0x3b, 0x2a)
OX_LT  = RGBColor(0xf7, 0xe9, 0xe5)
DARK   = RGBColor(0x1f, 0x1d, 0x18)
WHITE  = RGBColor(0xff, 0xff, 0xff)
DIM    = RGBColor(0xa8, 0xa2, 0x96)

# Standing rule for RCN documents: the attribution travels on the artifact,
# because decks get forwarded away from the conversation that produced them.
ATTRIB = 'Marc Pierson and Claude Opus 5  ·  September 2026'


prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(bgc=PAPER):
    sl = prs.slides.add_slide(BLANK)
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = bgc
    return sl


def rect(sl, x, y, w, h, fill=WHITE, line=None, lw=1.25):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def txt(sl, text, x, y, w, h, size=20, bold=False, italic=False,
        color=INK, align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return tb


def bullets(sl, items, x, y, w, h, size=17, color=INK2, gap=11, bullet='—'):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        bold = False
        if isinstance(it, tuple):
            it, bold = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.18
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = (bullet + '  ' if bullet else '') + it
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = INK if bold else color
        r.font.name = 'Calibri'
    return tb


def kicker(sl, text, color=ACC):
    txt(sl, text.upper(), 0.85, 0.55, 11, 0.32, size=11, bold=True, color=color)


def title(sl, text, y=0.95, size=34, color=INK, w=11.6):
    txt(sl, text, 0.85, y, w, 1.3, size=size, bold=True, color=color, spacing=0.95)


def footer(sl, n, dark=False):
    c = DIM if dark else INK3
    txt(sl, 'RCN method  ·  Kitchen Table Exercise', 0.85, 6.95, 7, 0.3, size=9, color=c)
    txt(sl, str(n), 12.1, 6.95, 0.4, 0.3, size=9, color=c, align=PP_ALIGN.RIGHT)


# -- Text fitting. PowerPoint does not shrink text to fit; it spills past the
#    bottom edge and the spill is invisible in the XML, so every box holding
#    variable prose measures first. Calibri averages ~0.505 em per character.
CHAR_EM = 0.505


def _lines(text, w_in, size_pt):
    cpl = max(8, int(w_in * 72 / (size_pt * CHAR_EM)))
    return sum(max(1, -(-len(p) // cpl)) for p in text.split('\n'))


def _text_h(text, w_in, size_pt, spacing=1.12):
    return _lines(text, w_in, size_pt) * size_pt * spacing * 1.02 / 72


def card(sl, x, y, w, h, head, body, headsize=15, bodysize=12.5, fill=WHITE, edge=RULE):
    tw = w - 0.5
    while bodysize > 9.0 and 0.68 + _text_h(body, tw, bodysize) + 0.12 > h:
        bodysize -= 0.5
    rect(sl, x, y, w, h, fill, edge)
    txt(sl, head, x + 0.28, y + 0.18, tw, 0.4, size=headsize, bold=True, color=INK)
    txt(sl, body, x + 0.28, y + 0.68, tw, h - 0.86, size=bodysize, color=INK2, spacing=1.12)


def numdisc(sl, cx, cy, d, n, fill=ACC_LT, ring=ACC, tc=ACC, size=17):
    sh = sl.shapes.add_shape(9, Inches(cx - d/2), Inches(cy - d/2), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = ring; sh.line.width = Pt(1.1)
    sh.shadow.inherit = False
    txt(sl, str(n), cx - d/2, cy - 0.155, d, 0.34, size=size, bold=True,
        color=tc, align=PP_ALIGN.CENTER)


n = 0
def nxt():
    global n
    n += 1
    return n


# ── 1  Title ────────────────────────────────────────────────────────────────
sl = slide(DARK)
txt(sl, 'AN RCN METHOD', 0.85, 1.75, 11, 0.32, size=12, bold=True, color=DIM)
txt(sl, 'The Kitchen Table\nExercise', 0.85, 2.3, 11, 2.0, size=54, bold=True,
    color=WHITE, spacing=0.95)
txt(sl, 'One hour. Paper and a pencil. Nobody helps you.',
    0.85, 4.45, 11, 0.5, size=21, color=PALE)
rect(sl, 0.85, 5.25, 3.4, 0.02, DIM)
txt(sl, 'Eleven steps, then an hour of looking back at them. At the end you will '
        'have made a clear sparse visual story about something you care about, out '
        'of nothing but your own mind and your own experience.',
    0.85, 5.55, 8.6, 1.1, size=13.5, color=DIM, spacing=1.2)
txt(sl, ATTRIB, 0.85, 6.85, 8.6, 0.32, size=11, color=DIM)

# ── 2  The promise ──────────────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'What happens')
title(sl, 'In about an hour you will have made a diagram\nyou can read out loud', size=31)
bullets(sl, [
    'You will choose what it is about. Nobody else will.',
    'You will not be taught anything, and nothing will be marked.',
    'No example will be shown to you, on purpose.',
    'Then you will spend about the same time again writing down how each step felt.',
], 0.85, 3.0, 11.3, 2.4, size=17.5)
rect(sl, 0.85, 5.75, 11.6, 0.95, ACC_LT, ACC)
txt(sl, 'The second hour is where most of the learning is. Do not skip it.',
    1.15, 6.0, 11, 0.5, size=16, bold=True, color=ACC)
footer(sl, nn)

# ── 3  What you need ────────────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'Before you begin')
title(sl, 'What you need')
for i, (h, b) in enumerate([
    ('An hour, and then another',
     'One hour for the drawing, and roughly the same again for the looking back. '
     'They do not have to be the same day, but do not leave a long gap.'),
    ('Paper, a pencil, an eraser',
     'Three or four clean sheets. The eraser matters more than you would think. '
     'No software, no licence, nothing to install.'),
    ('Somewhere you will not be interrupted',
     'Alone. If you are doing this in a room with other people, you are still '
     'working on your own paper and nobody is looking at it.'),
]):
    card(sl, 0.85 + i * 3.95, 2.75, 3.65, 2.5, h, b)
txt(sl, 'If you would rather keep your notes on paper than on a screen, print the '
        'Notes Sheet before you start.',
    0.85, 5.6, 11.3, 0.5, size=14, italic=True, color=INK2)
footer(sl, nn)

# ── 4  Three rules ──────────────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'Before you begin')
title(sl, 'Three rules')
for i, (h, b) in enumerate([
    ('Do it alone',
     'Even in a full room. Your paper is yours and nobody reads it unless you '
     'decide to show them.'),
    ('Do not read ahead',
     'One step at a time, in order. The exercise stops working if you skip '
     'forward, because you start aiming at what is coming.'),
    ('Nothing here will help you',
     'There is no worked example and no answer. If you find yourself wanting '
     'someone to check, notice that, and keep going.'),
]):
    y = 2.6 + i * 1.32
    numdisc(sl, 1.2, y + 0.45, 0.62, i + 1)
    txt(sl, h, 1.85, y + 0.06, 3.6, 0.45, size=19, bold=True, color=INK)
    txt(sl, b, 5.5, y + 0.02, 7.0, 1.0, size=14, color=INK2, spacing=1.15)
rect(sl, 0.85, 6.5, 11.6, 0.02, RULE)
footer(sl, nn)

# ── 5  The declaration ──────────────────────────────────────────────────────
sl = slide(DARK); nn = nxt()
txt(sl, 'SAID PLAINLY, IN ADVANCE', 0.85, 1.5, 11, 0.32, size=12, bold=True, color=OX)
txt(sl, 'This deck will not\nhelp you.', 0.85, 2.05, 11, 1.9, size=46, bold=True,
    color=WHITE, spacing=0.98)
txt(sl, 'It will give you one instruction at a time and nothing else. It will not '
        'show you what a good answer looks like, tell you whether yours is any '
        'good, or explain what a step was for until you have done all eleven.',
    0.85, 4.15, 8.3, 1.3, size=16, color=PALE, spacing=1.25)
txt(sl, 'That is the method, not a shortage of material. The reason is on the last '
        'few slides, where it will make sense.',
    0.85, 5.55, 8.3, 0.8, size=14, italic=True, color=DIM, spacing=1.2)
rect(sl, 9.6, 2.05, 2.85, 3.05, None, OX, 1.6)
txt(sl, 'Some steps\nmay feel\nuncomfortable.\n\nThat is common,\nand it is not a\nsign that you\nare doing it\nwrong.',
    9.9, 2.4, 2.4, 3.3, size=14.5, color=WHITE, spacing=1.3)
footer(sl, nn, dark=True)

# ── 6  Start ────────────────────────────────────────────────────────────────
sl = slide(); nn = nxt()
txt(sl, 'Turn the page when you are ready.', 0.85, 3.0, 11.6, 0.8, size=30,
    bold=True, color=INK, align=PP_ALIGN.CENTER)
txt(sl, 'From here on, one step per slide. Do the step, then advance.',
    0.85, 3.95, 11.6, 0.5, size=16, color=INK2, align=PP_ALIGN.CENTER)
footer(sl, nn)


# ── 7-17  The eleven steps — instruction only, no example, no decoration ────
STEPS = [
    ('Think of some area of present concern and care.',
     'It does not matter what the area is, only that it is really meaningful to '
     'you personally.\n\nTake your time. Be sure it is a big deal to you.'),
    ('Make a list of anything that comes to mind when you think about that topic.',
     'It is a list, not an essay. Feel free to list anything — an idea, a person, '
     'a place, a thing, any thing — that comes to your mind.\n\n'
     'Be sure there are at least fifteen items. Do not go further than what can be '
     'written on a single sheet of paper.\n\n'
     'Take your time. Take a break if you like.'),
    ('Choose about eight of the items from your list. You choose.',
     'Write each of these onto another clean sheet of paper. Spread them out, '
     'leaving lots of white space.\n\n'
     'Enclose each of the items separately — some sort of border around each one.'),
    ('Draw lines connecting some of the items from Step 3.',
     'Do this in ways that make sense to you.\n\nThe connections must matter to you.'),
    ('Label the connecting lines.',
     'What does the line mean?'),
    ('Add arrowheads to the lines.', ''),
    ('If you have more than one arrowhead on a line, erase one of the two.',
     'Draw another line to hold the extra arrowhead.'),
    ('Make sure every label on a line is a verb which conveys action or change.',
     'Typically a word ending in "-ing".\n\n'
     'You will probably need to spend some time on this step, and to erase some of '
     'the initial labels.\n\n'
     'By the way, you can change the labels on the items themselves any time you '
     'like. This is your graph, your drawing, about something that you care about.'),
    ('Read the graph. Out loud. Say only what is written on the page.',
     'I promise you that you can do this if you have used verbs on the connecting '
     'lines. If you have not done that, you will find the verbs when you try to '
     'read the graph.'),
    ('Now recall how you felt when you were asked to do Step 1.',
     'Really get back in touch with what was going on for you in that moment. Say it '
     'out loud. Listen to what you say. Write down what you have just heard yourself say.\n\n'
     'Then do the same for Step 2, and Step 3, and every step through Step 9.\n\n'
     'It typically takes as much time and thought as it took to do each step in the '
     'first place. This is where the big learning is.'),
    ('Make sure every step has a note saying how it felt and what you understood there.',
     'Then write down your answers to these three.\n\n'
     'What have you learned?\n\n'
     'What could you do with this little graph?\n\n'
     'Was this interesting? Fun? Useful?'),
]
for i, (head, body) in enumerate(STEPS):
    sl = slide(); nn = nxt()
    s = i + 1
    txt(sl, str(s), 9.9, 0.7, 2.6, 2.6, size=150, bold=True, color=PALE,
        align=PP_ALIGN.RIGHT)
    txt(sl, 'STEP ' + str(s), 0.85, 1.05, 6, 0.32, size=12, bold=True, color=ACC)
    hsize = 34
    while hsize > 22 and _text_h(head, 8.4, hsize, 1.0) > 1.9:
        hsize -= 1
    txt(sl, head, 0.85, 1.55, 8.4, 2.0, size=hsize, bold=True, color=INK, spacing=1.0)
    if body:
        bsize = 18
        while bsize > 12 and _text_h(body, 8.4, bsize, 1.3) > 2.6:
            bsize -= 0.5
        txt(sl, body, 0.85, 3.75, 8.4, 2.7, size=bsize, color=INK2, spacing=1.3)
    if s == 10:
        rect(sl, 0.85, 6.45, 8.4, 0.02, RULE)
        txt(sl, 'Take a few minutes first. Relax.', 0.85, 6.55, 8, 0.35,
            size=13, italic=True, color=INK3)
    footer(sl, nn)


# ── 18  The fold ────────────────────────────────────────────────────────────
sl = slide(DARK); nn = nxt()
txt(sl, 'STOP', 0.85, 1.6, 11.6, 1.2, size=58, bold=True, color=OX, align=PP_ALIGN.CENTER)
txt(sl, 'Everything after this slide assumes you have finished all eleven steps '
        'and written your notes.',
    2.4, 3.1, 8.5, 0.9, size=21, color=WHITE, align=PP_ALIGN.CENTER, spacing=1.2)
txt(sl, 'Reading it first will not make the exercise easier. It will make it not '
        'work, because you will spend the hour producing what you have been told '
        'to expect.',
    2.9, 4.3, 7.5, 1.2, size=15, color=DIM, align=PP_ALIGN.CENTER, spacing=1.25)
footer(sl, nn, dark=True)

# ── 19  The nine skills ─────────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'After you have finished')
title(sl, 'What you just did')
txt(sl, 'The list that was kept out of sight until now.', 0.85, 2.15, 11, 0.4,
    size=15, color=INK2)
NINE = [('Focusing', ''), ('Listing', 'looking about, imagining, naming, limiting'),
        ('Culling', 'narrowing, selecting'), ('Arraying', 'spatial organising'),
        ('Connecting', 'formally relating'), ('Assigning direction', 'vectoring'),
        ('Labelling', 'clarifying'), ('Narrating', 'telling the story'),
        ('Reflecting and synthesising', '')]
for i, (h, b) in enumerate(NINE):
    col, row = i % 3, i // 3
    x, y = 0.85 + col * 3.95, 2.7 + row * 1.1
    numdisc(sl, x + 0.26, y + 0.3, 0.5, i + 1, size=13)
    txt(sl, h, x + 0.62, y + 0.04, 3.2, 0.36, size=15, bold=True, color=INK)
    if b:
        txt(sl, b, x + 0.62, y + 0.42, 3.2, 0.5, size=11.5, color=INK3, spacing=1.1)
rect(sl, 0.85, 6.0, 11.6, 0.8, ACC_LT, ACC)
txt(sl, 'Every one of these is something you were already doing before anybody '
        'taught you anything. The exercise did not install a skill. It named one you had.',
    1.15, 6.18, 11.0, 0.5, size=13.5, bold=True, color=ACC, spacing=1.15)
footer(sl, nn)

# ── 20  Why it was hidden ───────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'After you have finished')
title(sl, 'Why that list was hidden until now')
txt(sl, 'Being told in advance what you are about to demonstrate turns it into a '
        'target. You would have spent the hour aiming at the list instead of '
        'thinking about your own concern.',
    0.85, 2.5, 8.6, 1.4, size=20, color=INK2, spacing=1.3)
txt(sl, 'It is the same reason there was no example. A worked example is a hint '
        'about what a good answer looks like, which puts a judge back in the room '
        'in written form.',
    0.85, 4.15, 8.6, 1.4, size=20, color=INK2, spacing=1.3)
rect(sl, 0.85, 5.95, 11.6, 0.02, RULE)
txt(sl, 'This is also why the person whose kitchen table it was did not look at '
        'anybody\'s paper.', 0.85, 6.1, 11.3, 0.5, size=14, italic=True, color=INK3)
footer(sl, nn)

# ── 21  Why some of it was hard ─────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'After you have finished', OX)
title(sl, 'One account of why some of it was hard')
rect(sl, 0.85, 2.15, 11.6, 0.95, OX_LT, OX)
txt(sl, 'This is a pattern observed across the people who have done this. It is not '
        'a diagnosis of you. Your own notes are better evidence about your hour than '
        'this slide is — if they disagree, trust your notes.',
    1.15, 2.33, 11.0, 0.7, size=14, color=OX, spacing=1.2)
txt(sl, 'The difficulty does not fall evenly. It concentrates wherever nobody is '
        'telling you whether you are right.',
    0.85, 3.45, 11.3, 0.9, size=24, bold=True, color=INK, spacing=1.15)
bullets(sl, [
    'Most of us learned to work in places where somebody else says whether it counts.',
    'This exercise removes that at every single step, and what people feel is the removal.',
    'The useful sentence is not "I am bad at this". It is "I was trained to need something that is not in this room".',
    'It is adaptive, not damage. It was a reasonable strategy where it was learned. It does not transfer here.',
], 0.85, 4.5, 11.3, 2.2, size=15.5)
footer(sl, nn)

# ── 22  Step by step ────────────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'After you have finished', OX)
title(sl, 'Where people tend to feel it', size=32)
PAIRS = [
    ('Step 1', 'Many people have not been asked in years what they actually care about — and whatever you choose exposes something.'),
    ('Step 2', 'The floor of fifteen pushes past where you would have stopped. Items eight to fifteen are the ones you were not going to admit to.'),
    ('Step 3', 'You select without being asked to justify it, and can spend the whole step bracing for a question that never comes.'),
    ('Step 5', 'Naming a relationship is where vocabulary insecurity lives, in anyone who has been corrected on their words.'),
    ('Step 8', 'Erasing has usually meant you were wrong. Here it means the work is progressing. Understanding that is easy; feeling it is not.'),
    ('Step 9', 'Reading your own work aloud is a familiar place to be humiliated — and now there is nobody to say whether it was any good.'),
]
for i, (h, b) in enumerate(PAIRS):
    col, row = i % 2, i // 2
    x, y = 0.85 + col * 5.95, 2.25 + row * 1.5
    card(sl, x, y, 5.65, 1.35, h, b, headsize=14, bodysize=12)
footer(sl, nn)

# ── 23  The range so far ────────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'After you have finished')
title(sl, 'The range so far, which runs backwards')
rect(sl, 0.85, 2.4, 5.5, 2.0, WHITE, RULE)
txt(sl, 'Found none of it hard', 1.15, 2.6, 5.0, 0.4, size=14, bold=True, color=ACC)
txt(sl, 'A student who had\nleft school', 1.15, 3.05, 5.0, 1.1, size=23, bold=True,
    color=INK, spacing=1.05)
rect(sl, 6.95, 2.4, 5.5, 2.0, WHITE, RULE)
txt(sl, 'Found every step very hard', 7.25, 2.6, 5.0, 0.4, size=14, bold=True, color=OX)
txt(sl, 'A lawyer', 7.25, 3.05, 5.0, 1.1, size=23, bold=True, color=INK, spacing=1.05)
txt(sl, 'Everybody else fell somewhere between the two.',
    0.85, 4.6, 11.3, 0.45, size=17, color=INK2, align=PP_ALIGN.CENTER)
txt(sl, 'That ordering is the reverse of what most schools would predict about those '
        'two people. It is what you would expect if the difficulty tracks years of '
        'successful submission to judgement rather than years of education.',
    0.85, 5.25, 11.3, 1.0, size=16, color=INK, spacing=1.25)
txt(sl, 'Fewer than twenty people, records kept informally or not at all. A direction, '
        'not a measurement.', 0.85, 6.45, 11.3, 0.4, size=13, italic=True, color=INK3)
footer(sl, nn)

# ── 24  What to do with it ──────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'After you have finished')
title(sl, 'What to do with your drawing')
for i, (h, b) in enumerate([
    ('Read it aloud to one person',
     'Do not hand it over and ask what they think. Read it to them. The difference '
     'matters and you will feel it.'),
    ('Do not tidy it up',
     'The wobbly version is the one you made. A neat copy is a different object and '
     'usually a worse one.'),
    ('Draw it again in a month',
     'From a blank sheet, not by editing this one. What changes between the two is '
     'the finding.'),
    ('Keep your notes',
     'Do this again on another topic and the two sets of notes side by side will say '
     'more than either does alone.'),
]):
    col, row = i % 2, i // 2
    card(sl, 0.85 + col * 5.95, 2.5 + row * 2.0, 5.65, 1.8, h, b, headsize=17, bodysize=13.5)
footer(sl, nn)

# ── 25  Hosting: the three components ───────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'Hosting it for other people')
title(sl, 'Running it is easy. Not helping is hard.')
for i, (h, b) in enumerate([
    ('Declare it in advance',
     'Say plainly, before they start, that you will not look at their work and will '
     'not help. Silence that has not been declared is not a withdrawal of judgement. '
     'It is surveillance, and it is worse than marking.'),
    ('Demonstrate it',
     'Do your own unrelated work while they draw. Not sitting attentively, not '
     'available if needed. Being visibly occupied is the proof that the declaration '
     'was true.'),
    ('Hold it',
     'They will ask. They will ask repeatedly, and some will ask in ways that are hard '
     'to refuse. Every refusal re-confirms the rule.'),
]):
    card(sl, 0.85 + i * 3.95, 2.65, 3.65, 2.55, h, b, headsize=17, bodysize=12.5)
rect(sl, 0.85, 5.55, 11.6, 1.05, ACC_LT, ACC)
txt(sl, 'Be present the whole time. This is not leaving them alone in a room. Alone, '
        'there is no authority to have a different kind of relationship with. At the '
        'table there is one, and it declines the role.',
    1.15, 5.75, 11.0, 0.7, size=14.5, color=ACC, spacing=1.2)
footer(sl, nn)

# ── 26  Hosting: the hint ───────────────────────────────────────────────────
sl = slide(DARK); nn = nxt()
txt(sl, 'HOSTING IT FOR OTHER PEOPLE', 0.85, 1.3, 11, 0.32, size=12, bold=True, color=OX)
txt(sl, 'A hint is worse than\na wrong answer.', 0.85, 1.85, 11, 1.7, size=42,
    bold=True, color=WHITE, spacing=0.98)
txt(sl, 'Give one hint and the protocol is over. The moment you show that you could '
        'have helped, your silence stops being a rule and becomes a withholding — and '
        'the whole hour converts backwards into a test with somebody in the room who '
        'knows the answer. From then on they are working out what you want.',
    0.85, 3.85, 7.6, 1.8, size=16, color=PALE, spacing=1.28)
rect(sl, 8.9, 1.85, 3.55, 3.75, None, DIM, 1.2)
txt(sl, 'THE LINE', 9.2, 2.1, 3.0, 0.3, size=11, bold=True, color=OX)
txt(sl, 'Anything said to everybody before Step 1 is framing, and it is allowed.\n\n'
        'Anything said to one person during the steps is help, and it is not.\n\n'
        'So a wall that lots of people hit gets fixed in the opening framing for '
        'everyone — never in a whisper to that person.',
    9.2, 2.5, 3.0, 3.4, size=12.5, color=PALE, spacing=1.25)
footer(sl, nn, dark=True)

# ── 27  Hosting: what you may say ───────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'Hosting it for other people')
title(sl, 'What you may say, and what you may not')
rect(sl, 0.85, 2.4, 5.65, 3.5, WHITE, ACC)
txt(sl, 'MAY SAY', 1.15, 2.6, 5.0, 0.3, size=11, bold=True, color=ACC)
bullets(sl, ['Take your time.',
             'There is no right answer to this one.',
             'Nobody will see this unless you decide to show them.',
             'Keep going.',
             'Yes, that is the whole instruction.',
             'I am not going to help you with that, and that is on purpose.'],
        1.15, 3.0, 5.05, 2.8, size=13.5, gap=7)
rect(sl, 6.8, 2.4, 5.65, 3.5, WHITE, OX)
txt(sl, 'MAY NOT SAY', 7.1, 2.6, 5.0, 0.3, size=11, bold=True, color=OX)
bullets(sl, ['Anything that evaluates.',
             'Anything that suggests.',
             'Any example, of anything, ever.',
             'Anything beginning "you could try".',
             'Any answer to "is this right".',
             'Any face at their paper — which is why you are not looking at it.'],
        7.1, 3.0, 5.05, 2.8, size=13.5, gap=7)
txt(sl, 'Their asking is not an interruption of the exercise. It is the most useful '
        'observation in it — and it belongs to the method, never to the person. A '
        'help-request count with a name on it is a score.',
    0.85, 6.1, 11.6, 0.7, size=14, italic=True, color=INK2, spacing=1.2)
footer(sl, nn)

# ── 28  The group version ───────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'Hosting it for other people')
title(sl, 'With a group: separately first, then the merge')
txt(sl, 'Everyone does their own, separately and simultaneously, in the same room. '
        'The group work starts only when each person is holding a finished drawing '
        'of their own.',
    0.85, 2.35, 11.3, 0.9, size=17, color=INK2, spacing=1.2)
for i, (h, b) in enumerate([
    ('Expect the merge to be slow',
     'Five students took six weekly sessions with sticky notes. The work was almost '
     'entirely about language — renaming, arguing about what words meant, combining '
     'near-synonyms. Any plan that treats the merge as a quick step after the fun '
     'part has it backwards.'),
    ('The group rule: do not resolve',
     'Alone you withhold the urge to teach. Here you withhold the urge to arbitrate. '
     'A facilitator\'s ruling does not just settle a question, it takes over the '
     'language — and the merge is about the language. Let a disagreement about a '
     'word stay open.'),
    ('Leave the seams',
     'The merged diagram will have the same idea in two places, because two people '
     'raised it and nobody had authority to rule. Leave them. They record where the '
     'group\'s language has not converged, which beats a tidy diagram.'),
]):
    card(sl, 0.85 + i * 3.95, 3.45, 3.65, 2.65, h, b, headsize=15.5, bodysize=12)
txt(sl, 'Encouragement is allowed because it carries no propositional content. '
        '"Keep going" says nothing about the map. "I think those two are the same '
        'thing" says everything about it.',
    0.85, 6.15, 11.6, 0.6, size=14, italic=True, color=INK2, spacing=1.2)
footer(sl, nn)

# ── 29  Where next ──────────────────────────────────────────────────────────
sl = slide(); nn = nxt()
kicker(sl, 'One more move, and it is worth the work')
title(sl, 'Turning your drawing into a causal loop diagram')
bullets(sl, [
    'Keep only the nodes that can be more or less — quantities that go up and down.',
    'Restate each verb as either same direction or opposite direction.',
    'Then look for a circle. The circles are feedback, which is where system behaviour lives.',
], 0.85, 2.5, 11.3, 1.9, size=18)
for i, (h, b) in enumerate([
    ('The question that does the work',
     '"What is it about this that matters here?" Forget for now whether you could '
     'measure it. The question is whether it matters.'),
    ('Keep the ones that refuse',
     'A node that will not become a variable is usually a thing rather than a quantity, '
     'or an abstraction hiding the quantity that matters. That list is the finding.'),
]):
    card(sl, 0.85 + i * 5.95, 4.5, 5.65, 1.85, h, b, headsize=16, bodysize=13)
footer(sl, nn)

# ── 30  Where it came from ──────────────────────────────────────────────────
sl = slide(DARK); nn = nxt()
txt(sl, 'WHERE THIS CAME FROM', 0.85, 1.35, 11, 0.32, size=12, bold=True, color=DIM)
txt(sl, 'Ferndale, Washington, 2019', 0.85, 1.85, 11, 0.9, size=36, bold=True, color=WHITE)
txt(sl, 'Marc Pierson developed the eleven steps and published them as Systems '
        'Certification. They were used with a cohort of students who had left school. '
        'One of them did it first. Then the cohort\'s facilitator did it, then her '
        'administrator, each alone at the same table — and then four more students at once.',
    0.85, 3.0, 7.7, 1.8, size=16, color=PALE, spacing=1.3)
txt(sl, 'The order of that spread — outward from the most disaffected student rather '
        'than down from the most senior adult — was not a gesture. He was the one who '
        'found it easiest.',
    0.85, 4.95, 7.7, 1.1, size=15, italic=True, color=DIM, spacing=1.3)
rect(sl, 9.0, 3.0, 3.45, 3.05, None, OX, 1.4)
txt(sl, 'The reflection steps\nwere not in the first\nversion.\n\nThey were added\nbecause the fear was\nvisible and needed\nsomewhere to go.',
    9.3, 3.3, 2.9, 2.6, size=14.5, color=WHITE, spacing=1.3)
rect(sl, 0.85, 6.35, 7.7, 0.015, DIM)
txt(sl, ATTRIB, 0.85, 6.5, 7.7, 0.32, size=11.5, color=DIM)
footer(sl, nn, dark=True)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                   'methods', 'kitchen-table', 'kitchen-table-exercise.pptx')
out = os.path.normpath(out)
prs.save(out)
print('Wrote', out, '·', len(prs.slides._sldIdLst), 'slides')
