"""
Generate sodoto-intro.docx from content matching sodoto-intro.html
Requires: pip install python-docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.25)
    section.right_margin  = Inches(1.25)

# ── Colour palette ───────────────────────────────────────────────────────────
GREEN  = RGBColor(0x2d, 0x6a, 0x4f)
BLUE   = RGBColor(0x1a, 0x56, 0xa4)
PURPLE = RGBColor(0x6b, 0x21, 0xa8)
AMBER  = RGBColor(0x92, 0x40, 0x0e)
RED    = RGBColor(0xb9, 0x1c, 0x1c)
GREY   = RGBColor(0x47, 0x55, 0x69)
BLACK  = RGBColor(0x1a, 0x1a, 0x1a)

# ── Helpers ──────────────────────────────────────────────────────────────────

def set_font(run, size=11, bold=False, italic=False, color=None, mono=False):
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    run.font.name  = 'Courier New' if mono else 'Calibri'

def heading(text, level=1, color=None):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = 'Calibri'
        if color:
            run.font.color.rgb = color
        if level == 1:
            run.font.size = Pt(22)
        elif level == 2:
            run.font.size = Pt(15)
            run.font.color.rgb = color or GREEN
        elif level == 3:
            run.font.size = Pt(12)
            run.font.color.rgb = color or GREEN
    return p

def para(text='', size=11, bold=False, italic=False, color=None, mono=False,
         space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic, color=color, mono=mono)
    return p

def bullet(text, level=0, size=11):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run, size=size)
    return p

def numbered(text, size=11):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, size=size)
    return p

def shade_cell(cell, hex_color):
    """Fill a table cell with a background colour."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def add_table(headers, rows, col_widths=None, header_fill='E8F5E9'):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    hdr_cells = t.rows[0].cells
    for i, h in enumerate(headers):
        shade_cell(hdr_cells[i], header_fill)
        run = hdr_cells[i].paragraphs[0].add_run(h)
        set_font(run, size=10, bold=True)

    # Data rows
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        fill  = 'FAFAF7' if ri % 2 == 1 else 'FFFFFF'
        for i, cell_text in enumerate(row):
            shade_cell(cells[i], fill)
            p = cells[i].paragraphs[0]
            if isinstance(cell_text, list):
                for fragment in cell_text:
                    txt, kw = (fragment[0], fragment[1]) if isinstance(fragment, tuple) else (fragment, {})
                    run = p.add_run(txt)
                    set_font(run, size=10, **kw)
            else:
                run = p.add_run(cell_text)
                set_font(run, size=10)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)

    doc.add_paragraph()   # spacer
    return t

def notice_box(label, text, label_color=None):
    """Render a callout block as a shaded paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(8)
    p.paragraph_format.left_indent  = Inches(0.2)
    r1 = p.add_run(label + '  ')
    set_font(r1, size=10, bold=True, color=label_color or RED)
    r2 = p.add_run(text)
    set_font(r2, size=10)

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT CONTENT
# ═══════════════════════════════════════════════════════════════════════════════

# ── Suite label ──────────────────────────────────────────────────────────────
p = doc.add_paragraph()
r = p.add_run('SODOTO DOCUMENTATION  ·  RELOCALIZE CREATIVITY NETWORK')
set_font(r, size=9, color=GREY)
p.paragraph_format.space_after = Pt(4)

# ── Title ────────────────────────────────────────────────────────────────────
heading('See One, Do One, Teach One', level=1)

p = doc.add_paragraph()
r = p.add_run(
    'A federated apprenticeship credentialing system for RCN practitioners. '
    'Skills are witnessed, practiced, and taught — then cryptographically certified '
    'by a Neighborhood Development Cooperative.'
)
set_font(r, size=12, italic=True, color=GREY)
p.paragraph_format.space_after = Pt(14)

# ── What SODOTO Is ───────────────────────────────────────────────────────────
heading('What SODOTO Is', level=2)

para(
    'SODOTO issues W3C Verifiable Credentials for practitioner skills. The credential '
    'documents not just that you have a skill, but how you acquired it: who showed you, '
    'who witnessed your practice, and who you subsequently taught. The full chain of '
    'transmission is embedded in the credential and cryptographically signed by the '
    'issuing NDC organization.'
)
para(
    'No central server is required to validate credentials. Anyone viewing a learner\'s '
    'FedWiki portfolio page can click Verify — the browser extracts the issuer\'s public '
    'key directly from the DID string and confirms the signature. No registration, no '
    'login, no third-party service.'
)
para(
    'Credentials are issued by NDC organizations, not individuals. The credentialing '
    'authority is the cooperative — persistent, institutional, not dependent on any one person.'
)
para(
    'The name comes from a medical training tradition: you observe a procedure done by an '
    'expert, you perform it yourself under supervision, and you teach it to the next learner. '
    'Each step is a gate. All three must pass before a credential is issued.'
)

# ── Roles ────────────────────────────────────────────────────────────────────
heading('Roles', level=2)

add_table(
    headers=['Role', 'Who', 'What They Do in SODOTO'],
    rows=[
        ['Learner',
         'Any RCN practitioner seeking a credential',
         'Goes through all three gates in sequence. Records their own assertion at each gate. '
         'May need multiple attempts at DoOne or TeachOne — all attempts are preserved.'],
        ['Mentor',
         'Someone who already holds the credential (or a founding cohort member)',
         'Demonstrates the skill at SeeOne. Witnesses the learner at DoOne and TeachOne. '
         'Makes a mentorAssertion at each gate. Their DID is embedded in the credential.'],
        ['Student',
         'The person the learner teaches at TeachOne',
         'Performs a DoOne session under the learner\'s teaching. Their DoOne FedWiki page '
         'is the evidence record for the learner\'s TeachOne. They begin their own credentialing path.'],
        ['NDC Coordinator',
         'Technical administrator at an issuing NDC',
         'After all three gates pass, runs the Veramo signing script to produce the credential '
         'JWT, commits the credential file to git, and creates or updates the learner\'s '
         'FedWiki portfolio page.'],
        ['Issuing NDC',
         'The cooperative organization (RCN, Columbia Valley NDC, The Fledge, etc.)',
         'Its Ed25519 key pair signs the JWT. Its DID appears in the credential as the issuer. '
         'Credentials are tied to the NDC\'s institutional identity, not any individual.'],
    ],
    col_widths=[1.2, 1.6, 3.2],
    header_fill='D1FAE5',
)

# ── Three-Gate Workflow ───────────────────────────────────────────────────────
heading('The Three-Gate Workflow', level=2)

para(
    'Every SODOTO credential passes through three sequential gates. Gates cannot be skipped. '
    'A gate may require multiple attempts — the record of each attempt, including failed ones, '
    'is preserved in the credential\'s history.'
)

# Gate descriptions as a table
add_table(
    headers=['Gate 1 — See One', 'Gate 2 — Do One', 'Gate 3 — Teach One'],
    rows=[[
        'The mentor demonstrates the skill in a real working context. The learner observes '
        'closely. Both record a brief assertion — the mentor that they showed the skill, the '
        'learner that they understood what they observed. The session is documented in a '
        'FedWiki narrative page.',
        'The learner performs the skill themselves. The mentor witnesses the full session. '
        'Both record assertions. If the mentor is not satisfied, the attempt is recorded as '
        'ended and another DoOne is scheduled. There is no penalty for multiple attempts — '
        'the history is evidence of rigor, not failure.',
        'The learner teaches a new student, with the mentor present to witness. Both the '
        'mentor and the student attest. The student\'s subsequent DoOne page serves as the '
        'evidence record — the learner cannot claim TeachOne without a real student who '
        'actually performed the skill.',
    ]],
    col_widths=[2.0, 2.0, 2.0],
    header_fill='DBEAFE',
)

heading('Assertion values recorded at each gate', level=3)

add_table(
    headers=['Gate', 'mentorAssertion', 'learnerAssertion', 'outcome'],
    rows=[
        ['SeeOne (pass)',      'showed',        'understood', 'complete'],
        ['DoOne (pass)',       'satisfied',     'has-done',   'complete'],
        ['DoOne (fail/pause)', 'not-satisfied', 'has-done',   'ended'],
        ['TeachOne (pass)',    'satisfied',     'has-taught', 'complete'],
    ],
    col_widths=[1.8, 1.6, 1.6, 1.0],
    header_fill='E8F0FE',
)

para(
    'Each gate also records the mentor\'s name and DID, a date, and a link to a FedWiki '
    'narrative page. The TeachOne gate adds the student\'s name, DID, portfolio slug, and '
    'the slug of the student\'s DoOne evidence page.',
    size=10, color=GREY
)

# ── System Architecture ───────────────────────────────────────────────────────
heading('System Architecture', level=2)

para(
    'The diagram below describes all components. The left side is the human process (people, '
    'gates). The center is the technical issuance engine. The right side is where credentials '
    'surface to learners and verifiers.'
)

# Architecture described as a structured table since SVG can't go in docx
arch_t = doc.add_table(rows=5, cols=3)
arch_t.style = 'Table Grid'
arch_t.alignment = WD_TABLE_ALIGNMENT.LEFT

arch_data = [
    ('PEOPLE', 'E8F5E9',
     'Learner (Practitioner)\nMentor (Credential Holder)\nNDC Coordinator (issues credential)',
     '', '',
     '', '', ''),

    ('GATE SEQUENCE', 'DBEAFE',
     'SEE ONE\nMentor demonstrates\nLearner observes\nboth assert · narrative page',
     '→',
     'DO ONE\nLearner performs\nMentor witnesses\nboth assert · may retry',
     '→',
     'TEACH ONE\nLearner teaches student\nMentor + student attest\nstudent\'s DoOne = evidence',
     ''),

    ('ISSUANCE ENGINE', 'EDE9FE',
     'Veramo agent  ·  Ed25519 keys (keys.json — gitignored, back up)\n'
     '5 NDC issuer DIDs  ·  signing script  ·  dids.json  ·  people.json\n'
     '→ output: Ed25519 JWT embedded in credential JSON',
     '', '', '', '', ''),

    ('CREDENTIAL STORE', 'F8FAFC',
     'veramo/credentials/*.json  (git committed)\n'
     'JWT embedded in each file\n'
     'dids.json · people.json (git committed)',
     '',
     'FEDWIKI + VERIFY',
     '',
     'Portfolio page · sodoto-badge plugin\n'
     'Verify button → Web Crypto API\n'
     'DID → public key → JWT check · entirely in-browser',
     ''),
]

labels    = ['PEOPLE', 'GATE SEQUENCE', 'ISSUANCE ENGINE', 'CREDENTIAL STORE / FEDWIKI']
fills     = ['D1FAE5', 'DBEAFE',        'EDE9FE',          'F1F5F9']
contents  = [
    'Learner (Practitioner)  ·  Mentor (Credential Holder)  ·  NDC Coordinator (issues credential)',
    'SEE ONE: Mentor demonstrates, Learner observes, both assert, narrative page  →  '
    'DO ONE: Learner performs, Mentor witnesses, both assert, may retry  →  '
    'TEACH ONE: Learner teaches student, Mentor + student attest, student\'s DoOne = evidence',
    'Veramo agent  ·  Ed25519 keys (keys.json — gitignored, back up)  ·  5 NDC issuer DIDs  '
    '·  signing script  ·  dids.json  ·  people.json  →  output: Ed25519 JWT embedded in credential JSON',
    'CREDENTIAL STORE: veramo/credentials/*.json (git committed), JWT embedded in each file, '
    'dids.json · people.json  |  '
    'FEDWIKI + VERIFY: Portfolio page · sodoto-badge plugin · Verify button → Web Crypto API · '
    'DID → public key → JWT check · entirely in-browser',
]

arch_t2 = doc.add_table(rows=4, cols=2)
arch_t2.style = 'Table Grid'
arch_t2.alignment = WD_TABLE_ALIGNMENT.LEFT

for i, (label, fill, content) in enumerate(zip(labels, fills, contents)):
    row = arch_t2.rows[i]
    shade_cell(row.cells[0], fill)
    r = row.cells[0].paragraphs[0].add_run(label)
    set_font(r, size=9, bold=True, color=GREY)
    row.cells[0].width = Inches(1.6)

    shade_cell(row.cells[1], 'FFFFFF')
    r2 = row.cells[1].paragraphs[0].add_run(content)
    set_font(r2, size=10)
    row.cells[1].width = Inches(4.4)

doc.add_paragraph()

# ── Issuing a Credential ─────────────────────────────────────────────────────
heading('Issuing a Credential — Coordinator Steps', level=2)

para(
    'The coordinator runs these steps after all three gates have passed and the FedWiki '
    'narrative pages for each gate are written.'
)

steps = [
    ('Confirm the skill is registered.',
     'Check that the skill has an established name in the registry. If it is a new skill, '
     'agree on the canonical name before proceeding — the name appears in the signed '
     'credential and cannot be changed after issuance.'),
    ('Confirm the learner has a DID.',
     'Look up the learner in ~/rcn/veramo/people.json. If they are not present, create an '
     'identity: node ~/rcn/veramo/create-did.js "Learner Name". Add the resulting DID and '
     'public key to both people.json and dids.json.'),
    ('Write the credential JSON file.',
     'Create a new file in ~/rcn/veramo/credentials/. Follow the naming convention: '
     'sodoto-[skill-slug]-[ndc]-[year]-[seq].json. Use an existing credential as the '
     'template. Populate all three gate sections with attempt arrays, mentor objects, '
     'dates, and FedWiki narrative page slugs. The jwt field stays empty at this stage.'),
    ('Sign the credential.',
     'Run the Veramo signing script with the issuing NDC\'s key. The script reads the '
     'credential JSON, signs it using Ed25519, and writes the JWT string back into the '
     'file\'s jwt field. (This step is currently manual and technical — a coordinator UI '
     'is planned.)'),
    ('Commit to git.',
     'git add veramo/credentials/ veramo/dids.json veramo/people.json and commit. '
     'Do NOT add keys.json — it is gitignored and must stay local.'),
    ('Add the badge to the learner\'s FedWiki portfolio.',
     'Open the learner\'s portfolio page in ~/.wiki/localhost/pages/. Add a sodoto-badge '
     'plugin item pointing to the credential filename. If no portfolio page exists for '
     'this learner, create one using Marc\'s portfolio as the template.'),
    ('Update the ledger.',
     'Add an entry to ~/.wiki/localhost/pages/rcn-sodoto-ledger recording the credential '
     'ID, skill, learner name, issuing NDC, and issuance date.'),
    ('Verify in the browser.',
     'Open the portfolio page in a browser and click the Verify button on the new badge. '
     'Confirm it returns a pass. If it fails, the JWT or the credential JSON structure '
     'has an error — check both before sharing.'),
]

for bold_text, rest in steps:
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(bold_text + '  ')
    set_font(r1, size=11, bold=True)
    r2 = p.add_run(rest)
    set_font(r2, size=11)

doc.add_paragraph()

notice_box(
    'Back up keys.json.',
    'The NDC\'s private keys live in ~/rcn/veramo/keys.json and are gitignored. '
    'If this file is lost, the NDC can no longer sign new credentials under those issuer DIDs. '
    'Existing credentials already signed remain valid — they only need the public key '
    '(embedded in the DID) to verify. Keep a secure offline backup.',
    label_color=AMBER,
)

# ── FedWiki Portfolio ─────────────────────────────────────────────────────────
heading("The Learner's FedWiki Portfolio", level=2)

para(
    'Each learner has a FedWiki page — for example, marc-pierson-sodoto-portfolio — that '
    'displays their credential badges. The page lives at ~/.wiki/localhost/pages/ and is '
    'served through the standard FedWiki server.'
)
para('Each badge on the portfolio page is rendered by the sodoto-badge plugin and displays:')

for item in [
    'The skill name and the issuing NDC\'s seal',
    'The three gate completion dates',
    'Mentor name(s), each linking to their own portfolio page',
    'The full attempt history for each gate (including any failed attempts)',
    'A Verify button that confirms the credential signature in-browser',
]:
    bullet(item)

para(
    'The badge plugin reads the credential\'s JWT from the page markup. The JWT is the '
    'credential — if the badge shows and the Verify button passes, the credential is '
    'authentic and unmodified.',
    space_before=6,
)

heading('Current portfolio pages', level=3)
add_table(
    headers=['Person', 'Page slug', 'Credentials'],
    rows=[
        ['Marc Pierson',  'marc-pierson-sodoto-portfolio',  '8 credentials (all skills issued to date)'],
        ['Kerry Turner',  'kerry-turner-sodoto-portfolio',   'Portfolio page exists; credentials to be populated'],
        ['Noah Williams', 'noah-williams-sodoto-portfolio',  'Portfolio page exists; credentials to be populated'],
    ],
    col_widths=[1.4, 2.4, 2.2],
    header_fill='D1FAE5',
)

# ── How Verification Works ────────────────────────────────────────────────────
heading('How Verification Works', level=2)

para(
    'The did:key format encodes the issuer\'s public key directly in the DID string. '
    'No server, no registry lookup, no external dependency. Anyone with the DID can '
    'recover the public key.'
)
para('When a user clicks Verify, the sodoto-badge plugin performs these steps entirely in the browser:')

for item in [
    'Extract the jwt string from the credential markup on the page',
    'Split the JWT into its three parts: header, payload, signature',
    'Decode the payload and read the iss field — this is the issuer\'s did:key string',
    'Decode the multibase suffix of the DID to recover the raw Ed25519 public key bytes',
    'Call crypto.subtle.verify() with those key bytes, the header.payload string, and the decoded signature',
    'Report ✓ Verified or ✗ Failed',
]:
    numbered(item)

doc.add_paragraph()
para(
    'This means credentials are permanently verifiable with no server, no expiry, and no '
    'reliance on any third party — as long as the DID string and the JWT are intact. '
    'The credential is self-contained.',
    space_before=4,
)

# ── Registered NDC Issuers ───────────────────────────────────────────────────
heading('Registered NDC Issuers', level=2)

para(
    'Five NDC organizations currently have Ed25519 key pairs and are registered in the '
    'badge plugin\'s ISSUER_REGISTRY with seal branding. The full DIDs are in '
    '~/rcn/veramo/dids.json.'
)

add_table(
    headers=['NDC', 'Location', 'Status'],
    rows=[
        ['ReLocalize Creativity Network', 'Bellingham WA',       'Active — 8 credentials issued'],
        ['Columbia Valley NDC',           'Whatcom County WA',   'Keys registered, no credentials yet'],
        ['The Fledge',                    'Lansing MI',          'Keys registered, no credentials yet'],
        ["Leo's",                         '(location TBD)',      'Placeholder'],
        ['Kula',                          '(location TBD)',      'Placeholder'],
    ],
    col_widths=[2.2, 1.6, 2.2],
    header_fill='D1FAE5',
)

# ── Skills Registry ───────────────────────────────────────────────────────────
heading('Skills Registry', level=2)

para('Green = credential issued.  Grey = planned, not yet issued.')

issued = [
    'Causal Loop Diagramming',
    'e-VSM Basic',
    'e-VSM Intermediate',
    'e-VSM Site Manager',
    'EIP Basic',
    'EIP Intermediate',
    'EIP Expert',
    'Stock and Flow Diagramming',
]
planned = [
    'Stock Flow Modeling',
    'System Dynamics Modeling',
    'VSM',
    'EIP Stage Sketching',
    'Six Context Questions',
    'Social Action Tetrahedron (PIOA)',
    'Social System Tetrahedron (BGTE)',
    'Cynefin',
    '15 Ps',
    'Process Mapping',
    'Object Process Methodology (OPM)',
    'Vester Sensitivity Model',
    'FedWiki',
    'Neo4j',
    'Persistent Syntegrating',
    'DSRP',
    'Ganz Organizing Story Sequence',
    'Campfire Conversation / Cave Drawing Set',
    'Conversation Taxonomy',
    'A3 Problem Solving',
]

# Two-column skill table
all_skills = [('✓  ' + s, True) for s in issued] + [('○  ' + s, False) for s in planned]
rows = []
for i in range(0, len(all_skills), 2):
    left  = all_skills[i]
    right = all_skills[i + 1] if i + 1 < len(all_skills) else ('', False)
    rows.append((left, right))

t = doc.add_table(rows=len(rows), cols=2)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, (left, right) in enumerate(rows):
    cells = t.rows[i].cells
    for ci, (text, is_issued) in enumerate([(left, left[1]), (right, right[1])]):
        shade_cell(cells[ci], 'F0FFF4' if text[1] else 'FAFAFA')
        r = cells[ci].paragraphs[0].add_run(text[0])
        set_font(r, size=10, color=GREEN if text[1] else GREY)
    cells[0].width = Inches(2.9)
    cells[1].width = Inches(2.9)
doc.add_paragraph()

notice_box(
    'Vester cluster note:',
    'The Vester Sensitivity Model cluster (nine interacting skills) will need a credential '
    'cluster design before individual credentials can be issued. The skills are deeply '
    'interdependent and a single linear credentialing path may not be appropriate.',
    label_color=AMBER,
)

# ── What Is Not Yet Built ─────────────────────────────────────────────────────
heading('What Is Not Yet Built', level=2)

para('The following items are designed, partially built, or simply needed — but not yet complete as of May 2026.')

gaps = [
    ('Credential issuance UI.',
     'Coordinators currently write credential JSON by hand and run CLI scripts. A '
     'browser-based issuance form is planned that would guide a coordinator through the '
     'gate fields and call the signing script automatically.'),
    ('DID onboarding for new practitioners.',
     'There is no defined self-service process for a new learner to get a DID. Currently '
     'the coordinator creates one via script and manually adds it to the registry files. '
     'A lightweight onboarding flow is needed.'),
    ('v0.4 credential format.',
     'Planned additions: a learner JWT (the learner cryptographically acknowledges each '
     'passing gate attempt), a student JWT (the student signs their affirmation at '
     'TeachOne), and a Neo4j debt record created automatically when TeachOne completes. '
     'None of these are in the current v0.3 format.'),
    ('Real gate histories on 7 credentials.',
     "Marc's CLD credential has full two-attempt histories. The e-VSM, EIP, and Stock "
     'and Flow credentials have single clean-pass placeholder histories. Real dates, '
     'mentor names, and narrative page links are needed when available.'),
    ('wiki-plugin-sodoto-badge source not committed to repo.',
     'The plugin source directory is installed in ~/.wiki/localhost/assets/ and '
     '~/node_modules/ but the source has not been committed to ~/rcn. It needs to be '
     'staged and committed so the repo is self-contained.'),
    ('CLI credential verification script.',
     'A coordinator tool that takes a signed credential JSON, extracts the signing '
     'payload, and confirms each Ed25519 proof against the issuer DID\'s public key — '
     'independent of the browser verify button. Useful for scripted audits.'),
    ('CfA-dSC integration.',
     'The Dyadic Smart Contract system (Conversations for Action + three-currency '
     'settlement) is a related but separate tool stack. A contract creator UI exists at '
     '~/rcn/tools/cfa-dsc-creator.html, Veramo DIDs are at ~/veramo-rcn/, and FedWiki '
     'party-view pages were working as of March 2026. The connection between CfA-dSC '
     'contracts and SODOTO credentials is designed but not implemented.'),
    ('cfa-dsc-schema.json and validator library not in repo.',
     'The JSON schema and TypeScript validator (25 tests) for the CfA-dSC system were '
     'built in a prior session but are not present on disk. Need reconstruction or '
     'location and committing to ~/rcn.'),
]

for bold_text, rest in gaps:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.15)
    p.paragraph_format.space_after  = Pt(8)
    r1 = p.add_run(bold_text + '  ')
    set_font(r1, size=11, bold=True, color=RED)
    r2 = p.add_run(rest)
    set_font(r2, size=11)

# ── Save ─────────────────────────────────────────────────────────────────────
out = '/Users/marcpierson/rcn/docs/sodoto-intro.docx'
doc.save(out)
print(f'Saved: {out}')
