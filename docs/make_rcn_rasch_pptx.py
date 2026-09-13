"""
Generate tools/rcn-rasch-intro.pptx — the RCN Rasch intro deck.

Companion to tools/rcn-rasch-intro.html and rcn-rasch-manual.html; same
argument, deck shape. Regenerate after editing:
    python3 docs/make_rcn_rasch_pptx.py
Requires: pip install python-pptx

AUDIENCE: the same measurement-minded engineers the SPC deck was written for,
plus a co-op board that will be handed a Wright map and asked to read it.
Leads with the "a total is not a measure" problem, shows the Wright map as
the product, and spends its weight on the three honest behaviours — misfit
named in words, extremes drawn without a number, thresholds checked before a
questionnaire is trusted.

The numbers on the case slides are REAL OUTPUT from the two demo datasets in
tools/rcn-rasch.html (Badges, Likert), taken from a headless run of the tool.
The Wright map slide is drawn with shapes from those same measures. If the
demo generators change, rerun both demos and change the data block to match.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ── Palette — steel blue persons, amber items, misfit red; matches the tool ──
ACC    = RGBColor(0x1e, 0x5f, 0x8c)   # steel blue — persons, this tool
ACC_LT = RGBColor(0xdb, 0xea, 0xf6)
ACC_DK = RGBColor(0x12, 0x3a, 0x5a)
ITEM   = RGBColor(0xb4, 0x53, 0x09)   # amber — items
ITEM_LT= RGBColor(0xfe, 0xf3, 0xc7)
MIS    = RGBColor(0xdc, 0x26, 0x26)   # misfit red
MIS_LT = RGBColor(0xfe, 0xf2, 0xf2)
GRN    = RGBColor(0x15, 0x80, 0x3d)
GRN_LT = RGBColor(0xdc, 0xfc, 0xe7)
AMB_LT = RGBColor(0xfe, 0xf9, 0xc3)
AMB_ED = RGBColor(0xfd, 0xe0, 0x47)
AMB_TX = RGBColor(0x85, 0x4d, 0x0e)
INK    = RGBColor(0x0f, 0x17, 0x2a)
INK2   = RGBColor(0x47, 0x55, 0x69)
INK3   = RGBColor(0x94, 0xa3, 0xb8)
SURF   = RGBColor(0xf8, 0xfa, 0xfc)
BORDER = RGBColor(0xe2, 0xe8, 0xf0)
RULER  = RGBColor(0x33, 0x41, 0x55)
WHITE  = RGBColor(0xff, 0xff, 0xff)
PALE   = RGBColor(0x7b, 0xa7, 0xc7)
PALE2  = RGBColor(0xc6, 0xdb, 0xea)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Primitives ──────────────────────────────────────────────────────────────

def slide(bgc=SURF):
    sl = prs.slides.add_slide(BLANK)
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = bgc
    return sl


def rect(sl, x, y, w, h, fill=WHITE, line=None, lw=1.25, dash=None):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = dash
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def oval(sl, cx, cy, d, fill, line=None, lw=1.2, dash=None):
    sh = sl.shapes.add_shape(9, Inches(cx - d/2), Inches(cy - d/2), Inches(d), Inches(d))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = dash
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def conn(sl, x1, y1, x2, y2, color, lw=1.2, dash=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    return c


def txt(sl, text, x, y, w, h, size=20, bold=False, italic=False,
        color=INK, align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return tb


def bullets(sl, items, x, y, w, h, size=17, color=INK2, gap=10, bullet='—'):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        bold = False
        if isinstance(it, tuple):
            it, bold = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = (bullet + '  ' if bullet else '') + it
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = INK if bold else color
        r.font.name = 'Calibri'
    return tb


def kicker(sl, text, color=ACC):
    txt(sl, text.upper(), 0.85, 0.55, 11, 0.32, size=11, bold=True, color=color)


def title(sl, text, y=0.95, size=34, color=INK, w=11.6):
    txt(sl, text, 0.85, y, w, 1.3, size=size, bold=True, color=color, spacing=0.95)


def footer(sl, n):
    txt(sl, 'RCN · Rasch', 0.85, 6.95, 6, 0.3, size=9, color=INK3)
    txt(sl, str(n), 12.1, 6.95, 0.4, 0.3, size=9, color=INK3, align=PP_ALIGN.RIGHT)


# ── Text fitting ────────────────────────────────────────────────────────────
# PowerPoint does not shrink text to fit a shape; it spills past the bottom
# edge and the spill is invisible in the XML. Every box holding variable
# prose measures it first. Calibri averages ~0.505 em per character.

CHAR_EM = 0.505
BOTTOM  = 6.75


def _lines(text, w_in, size_pt):
    cpl = max(8, int(w_in * 72 / (size_pt * CHAR_EM)))
    return sum(max(1, -(-len(p) // cpl)) for p in text.split('\n'))


def _text_h(text, w_in, size_pt, spacing=1.12):
    return _lines(text, w_in, size_pt) * size_pt * spacing * 1.02 / 72


def card(sl, x, y, w, h, head, body, accent=ACC, headsize=15, bodysize=12.5):
    tw = w - 0.5
    while bodysize > 9.0 and 0.72 + _text_h(body, tw, bodysize) + 0.12 > h:
        bodysize -= 0.5
    rect(sl, x, y, w, h, WHITE, BORDER)
    rect(sl, x, y, 0.055, h, accent)
    txt(sl, head, x + 0.28, y + 0.2, w - 0.5, 0.4, size=headsize, bold=True, color=INK)
    txt(sl, body, x + 0.28, y + 0.72, tw, h - 0.9, size=bodysize, color=INK2, spacing=1.1)


def banner(sl, x, y, w, h, head, body, fill, edge, headcolor, size=12.5):
    tw = w - 0.6
    need = lambda s: 0.58 + _text_h(body, tw, s) + 0.12
    room = max(h, BOTTOM - y)
    while size > 9.0 and need(size) > room:
        size -= 0.5
    h = max(h, min(need(size), room))
    rect(sl, x, y, w, h, fill, edge)
    txt(sl, head, x + 0.3, y + 0.15, tw, 0.36, size=15, bold=True, color=headcolor)
    txt(sl, body, x + 0.3, y + 0.58, tw, h - 0.72, size=size, color=INK2, spacing=1.12)
    return h


def label(sl, text, x, y, w, size, color, align=PP_ALIGN.LEFT, bold=True):
    """Tight label with zero insets, for annotating drawings."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return tb


def table(sl, x, y, colw, rows, size=12, header=True, rowh=0.36, numcols=()):
    """Simple ruled table drawn from text boxes, so widths and fonts are ours."""
    yy = y
    for ri, row in enumerate(rows):
        xx = x
        is_head = header and ri == 0
        for ci, cell in enumerate(row):
            al = PP_ALIGN.RIGHT if ci in numcols else PP_ALIGN.LEFT
            col = INK3 if is_head else (INK if ci == 0 else INK2)
            tb = sl.shapes.add_textbox(Inches(xx), Inches(yy), Inches(colw[ci]), Inches(rowh))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.06)
            tf.margin_top = tf.margin_bottom = Inches(0.03)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = al
            r = p.add_run()
            bold = is_head or ci == 0
            if isinstance(cell, tuple):
                cell, colr = cell
                col = colr
                bold = True
            r.text = str(cell)
            r.font.size = Pt(size - 2 if is_head else size)
            r.font.bold = bold
            r.font.color.rgb = col
            r.font.name = 'Calibri'
            xx += colw[ci]
        yy += rowh
        conn(sl, x, yy, x + sum(colw), yy, BORDER if ri else INK3, 0.75)
    return yy


# ── The data behind the case slides (real output of the tool's demos) ──────
# Badge demo: 24 persons × 12 skills, dichotomous, JMLE, 16 iterations.
PERSONS = [  # (name, measure) — omar is extreme (12/12), no measure
    ('ray', 4.09), ('sam', 4.09), ('tess', 4.09), ('lee', 2.80), ('nina', 2.80),
    ('pia', 1.84), ('quinn', 1.84), ('kim', 0.37), ('mo', 0.37), ('ivan', -0.25),
    ('chris', -0.82), ('eli', -0.82), ('fran', -0.82), ('gus', -0.82), ('hana', -0.82),
    ('jo', -0.82), ('ana', -1.37), ('kerry', -1.95), ('dorothy', -2.62), ('luis', -2.62),
    ('ben', -2.62), ('dee', -2.62), ('marc', -3.55)]
ITEMS = [  # (name, measure, flag) — flag is False, 'misfit' or 'provisional'
    ('stand up an NDC server', 4.27, 'provisional'), ('build a value network', 2.98, False),
    ('assign an IAD level', 1.91, False), ('label a value network arrow', 1.00, False),
    ('issue a signed badge', 0.59, False), ('fork from another site', -0.44, False),
    ('draw a simple graph', -1.06, False), ('edit a wiki page', -1.37, False),
    ('add a page to a lineup', -1.37, False), ('run a campfire conversation', -1.37, 'misfit'),
    ('publish a FedWiki page', -2.37, False), ('comment on a page', -2.76, False)]
BADGE_STATS = dict(rel=0.84, sep=2.32, strata=3.4, sd=2.31, rmse=0.92, irel=0.90, targ=-0.01)
# Likert demo: 140 residents × 7 statements, 1–5, Andrich, 28 iterations.
THRESH = [(1, 215, 21.9, None), (2, 276, 28.2, -1.81), (3, 95, 9.7, 0.63),
          (4, 205, 20.9, -0.31), (5, 189, 19.3, 1.48)]
LIKERT_STATS = dict(rel=0.88, sep=2.72, strata=4.0, isep=9.57)

n = 0
def nxt():
    global n
    n += 1
    return n


def wright_map(sl, x, y, w, h, persons, items, extreme_top=('omar',), lo=-4, hi=5):
    """Persons left, items right, one ruler. Drawn from measures, to scale."""
    cx = x + w * 0.46
    top, bot = y + 0.6, y + h - 0.15
    def Y(m): return top + (hi - m) / (hi - lo) * (bot - top)
    conn(sl, cx, top, cx, bot, RULER, 2.0)
    for t in range(lo, hi + 1):
        yy = Y(t)
        conn(sl, cx - 0.07, yy, cx + 0.07, yy, RULER, 1.2)
        label(sl, ('+' if t > 0 else '') + str(t), cx - 0.42, yy - 0.09, 0.32, 8, INK3, PP_ALIGN.RIGHT)
    label(sl, 'PERSONS — more able above', cx - 3.3, y, 3.0, 9, ACC, PP_ALIGN.RIGHT)
    label(sl, 'SKILLS — harder above', cx + 0.3, y, 3.0, 9, ITEM)
    label(sl, 'logits', cx - 0.3, bot + 0.02, 0.6, 8, INK3, PP_ALIGN.CENTER, bold=False)
    # persons, binned to a quarter logit like the tool
    bins = {}
    for name, m in persons:
        bins.setdefault(round(m / 0.25), []).append(name)
    for k, names in bins.items():
        yy = Y(k * 0.25)
        for j in range(len(names)):
            oval(sl, cx - 0.5 - j * 0.17, yy, 0.11, ACC)
        label(sl, ', '.join(names), cx - 0.64 - len(names) * 0.17 - 2.6, yy - 0.09, 2.6, 8, INK2,
              PP_ALIGN.RIGHT, bold=False)
    # items, nudged apart
    prev = -9
    for name, m, flag in sorted(items, key=lambda i: -i[1]):
        yy = max(Y(m), prev + 0.145)   # pack downward, tightly
        prev = yy
        mis = flag == 'misfit'
        if flag == 'provisional':
            rect(sl, cx + 0.22, yy - 0.065, 0.1, 0.13, None, ITEM, 1.2, MSO_LINE_DASH_STYLE.SQUARE_DOT)
            label(sl, name + '   provisional — 2 on the minority side', cx + 0.4, yy - 0.085, 3.4, 8, INK3, bold=False)
            continue
        rect(sl, cx + 0.22, yy - 0.065, 0.1, 0.13, MIS if mis else ITEM)
        label(sl, name + ('   misfit' if mis else ''), cx + 0.4, yy - 0.085, 3.2, 8,
              MIS if mis else INK2, bold=mis)
    # extremes above the top tick
    if extreme_top:
        ty = top - 0.2
        conn(sl, x + 0.2, ty + 0.1, x + w - 0.2, ty + 0.1, BORDER, 0.75, MSO_LINE_DASH_STYLE.DASH)
        for j, name in enumerate(extreme_top):
            oval(sl, cx - 0.5 - j * 0.17, ty, 0.11, None, ACC, 1.5, MSO_LINE_DASH_STYLE.SQUARE_DOT)
        label(sl, ', '.join(extreme_top) + ' — above every skill, no measure',
              cx - 0.64 - len(extreme_top) * 0.17 - 2.6, ty - 0.09, 2.6, 8, INK3, PP_ALIGN.RIGHT, bold=False)
    # means
    pm = sum(m for _, m in persons) / len(persons)
    conn(sl, x + 0.2, Y(pm), cx - 0.12, Y(pm), ACC, 0.75, MSO_LINE_DASH_STYLE.DASH)
    label(sl, 'mean person', x + 0.2, Y(pm) - 0.2, 1.2, 7.5, ACC)
    conn(sl, cx + 3.7, Y(0), x + w - 0.2, Y(0), ITEM, 0.75, MSO_LINE_DASH_STYLE.DASH)
    label(sl, 'mean skill = 0', x + w - 1.4, Y(0) - 0.2, 1.2, 7.5, ITEM, PP_ALIGN.RIGHT)


# ═════════════════════════════════════════════════════════════════════════════
# 1 — Title
# ═════════════════════════════════════════════════════════════════════════════
sl = slide(ACC_DK)
rect(sl, 0, 0, 13.33, 7.5, ACC_DK)
rect(sl, 0.85, 2.15, 1.6, 0.055, ITEM)
txt(sl, 'ReLocalize Creativity Network', 0.85, 1.55, 8, 0.35, size=12.5, bold=True, color=PALE)
txt(sl, 'RCN Rasch', 0.85, 2.5, 11, 1.1, size=48, bold=True, color=WHITE)
txt(sl, 'Why a total score is not a measurement, and what a Wright map adds',
    0.85, 3.65, 10.5, 0.5, size=19, color=PALE2)
txt(sl, 'Rasch measurement · after Rasch, Wright and Linacre · built for the SODOTO badge ledger',
    0.85, 4.25, 10.5, 0.4, size=13, italic=True, color=PALE)
txt(sl, 'September 2026', 0.85, 6.6, 4, 0.3, size=11, color=RGBColor(0x5d, 0x82, 0xa0))

# ═════════════════════════════════════════════════════════════════════════════
# 2 — You are already counting
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Where we are')
title(sl, 'You are already counting')
txt(sl, 'Every one of these produces a total. None of them produces a measure.',
    0.85, 2.15, 11.6, 0.5, size=19, color=INK2)
xs = [0.85, 3.85, 6.85, 9.85]
for x, (h, b) in zip(xs, [
        ('Badges held', 'SODOTO ledger — how many of the skills a person has been badged on'),
        ('Survey scores', 'seven statements, 1 to 5 each, added up to a number out of 35'),
        ('Checklist items met', 'a site audit — how many of the practices are in place'),
        ('Activation levels', 'the Patient Activation Measure, which came out of exactly this method')]):
    card(sl, x, 2.95, 2.75, 1.5, h, b, ACC, headsize=14, bodysize=11.5)
banner(sl, 0.85, 4.85, 11.6, 1.35,
       'A total counts rungs on a ladder nobody has measured',
       'Ten badges is more than two. But is it five times as capable? Which two? Is the step from the '
       'second badge to the third the same size as the step from the ninth to the tenth? A count '
       'cannot say. It assumes every item is equally hard and every step is the same size — and '
       'neither is ever checked.',
       ACC_LT, ACC, ACC_DK)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 3 — The line
# ═════════════════════════════════════════════════════════════════════════════
sl = slide(ACC_DK)
rect(sl, 0, 0, 13.33, 7.5, ACC_DK)
rect(sl, 0.85, 2.0, 1.6, 0.06, ITEM)
txt(sl, 'A total score tells you how many.', 0.85, 2.45, 12.0, 0.8, size=38, bold=True, color=WHITE)
txt(sl, 'It does not tell you how much.', 0.85, 3.35, 12.0, 0.8, size=38, bold=True, color=RGBColor(0xf5, 0xb0, 0x4a))
txt(sl, 'Measurement puts the people and the items on one ruler with known spacing. That is a different '
        'thing from adding up, and only one model makes adding up defensible — because it is the one '
        'in which the raw score is enough to place a person on the ruler.',
    0.85, 4.55, 10.5, 1.2, size=16, color=PALE2, spacing=1.2)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 4 — One ruler
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The model')
title(sl, 'One ruler, two kinds of thing on it')
txt(sl, 'Georg Rasch, 1960. The odds that a person succeeds at an item depend only on the distance between them.',
    0.85, 2.15, 11.6, 0.5, size=17, color=INK2)
rect(sl, 0.85, 2.85, 5.4, 1.15, WHITE, BORDER)
txt(sl, 'P(success)  =  e^(B − D) / (1 + e^(B − D))', 1.1, 3.0, 5.0, 0.45, size=19, bold=True, color=ACC_DK)
txt(sl, 'B = person ability     D = item difficulty     both in logits, one scale',
    1.1, 3.5, 5.0, 0.4, size=11.5, color=INK2)
table(sl, 6.7, 2.85, [2.2, 1.6, 2.0],
      [['Person minus item', 'Odds', 'Success'],
       ['+2 logits', '7 : 1', '88%'],
       ['+1 logit', 'e : 1', '73%'],
       ['0', 'even', '50%'],
       ['−1 logit', '1 : e', '27%'],
       ['−2 logits', '1 : 7', '12%']], size=12, rowh=0.34, numcols=(1, 2))
banner(sl, 0.85, 5.05, 11.6, 1.2,
       'What makes it a measurement rather than a score',
       'Persons and items land on the same ruler. A skill at +1.9 logits is not "hard" in the abstract; a '
       'person at +1.9 has even odds on it. That is a claim you can check in a room. And a person\'s raw '
       'score is a sufficient statistic for their measure — the one model in which adding up is honest, '
       'and it tells you the spacing of the rungs as it goes.',
       ACC_LT, ACC, ACC_DK)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 5 — The Wright map (drawn from the badge demo)
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The product')
title(sl, 'The Wright map', size=32)
txt(sl, 'The badge demo: 24 people, 12 SODOTO skills. Persons left, skills right, one ruler in logits.',
    0.85, 1.75, 8, 0.4, size=13, color=INK2)
rect(sl, 0.85, 2.2, 8.2, 4.55, WHITE, BORDER)
wright_map(sl, 0.95, 2.35, 8.0, 4.3, PERSONS, ITEMS)
card(sl, 9.35, 2.2, 3.1, 1.4, 'Read across, not down',
     'A skill level with a person is one they have even odds on. Skills far above everyone are out of reach; '
     'far below, already universal.', ACC, headsize=13, bodysize=10.5)
card(sl, 9.35, 3.75, 3.1, 1.4, 'Targeting −0.01',
     'Mean person against mean skill. The skills are aimed where the people are — a real instrument '
     'rarely manages that on its first run.', GRN, headsize=13, bodysize=10.5)
card(sl, 9.35, 5.3, 3.1, 1.45, 'Two things to see first',
     'Is there a skill near every cluster of people, or are there gaps on the ruler? And is anything red? A hollow bar is not red — it is too few responses to say.',
     ITEM, headsize=13, bodysize=10.5)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 6 — Misfit is information
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Honest behaviour 1 of 3')
title(sl, 'Misfit is information, not an error code', size=32)
txt(sl, 'The model predicts every cell. Where responses disagree with the predictions more than chance allows, '
        'it says so — per item, per person, in words.',
    0.85, 2.05, 11.6, 0.6, size=15, color=INK2)
table(sl, 0.85, 2.85, [3.4, 1.1, 1.1, 1.0, 1.0, 4.0],
      [['Skill', 'Badged', 'Measure', 'Infit', 'Outfit', 'Reading'],
       ['stand up an NDC server', '2 of 23', '+4.27', '1.36', ('2.31', INK3), ('Provisional — only 2 succeeded, too few to read fit', INK3)],
       ['assign an IAD level', '6 of 23', '+1.91', ('0.40', ITEM), ('0.18', ITEM), 'Too predictable — tracks its neighbours'],
       ['fork from another site', '12 of 23', '−0.44', '0.93', '0.84', 'Behaves as expected'],
       [('run a campfire conversation', MIS), '15 of 23', '−1.37', ('2.07', MIS), ('11.61', MIS), ('Badly noisy — measuring something else', MIS)],
       ['comment on a page', '19 of 23', '−2.76', '1.05', ('0.66', INK3), ('Provisional — only 4 failed, too few to read fit', INK3)]],
      size=11.5, rowh=0.36, numcols=(1, 2, 3, 4))
banner(sl, 0.85, 5.2, 11.6, 1.3,
       'Eleven of twelve behave or sit on too little evidence to say. One does not, and it is not close.',
       'Some of the least able people hold the campfire badge and some of the most able do not. The demo '
       'built it that way — running a campfire depends on temperament, not on the FedWiki-and-graph '
       'capability the other eleven share — and the model found it without being told. In a real ledger '
       'that is not a nuisance. It is the finding: this badge belongs to a different construct.',
       MIS_LT, RGBColor(0xfe, 0xca, 0xca), RGBColor(0x99, 0x1b, 0x1b))
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 7 — Infit and outfit
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Reading the fit')
title(sl, 'Two fit statistics, and why there are two')
card(sl, 0.85, 2.15, 5.6, 1.9, 'Infit — information-weighted',
     'Weights responses near a person\'s own level, where the model is least sure and the response '
     'carries most information. Reflects the item as a whole. An item with high infit means something '
     'different to different people.', ACC, headsize=15, bodysize=12.5)
card(sl, 6.85, 2.15, 5.6, 1.9, 'Outfit — unweighted',
     'Dominated by the few responses far from a person\'s level: the lucky guess on something far too '
     'hard, the careless miss on something far too easy. Ray holds 11 of 12 badges, misses the easy '
     'campfire one, and has an outfit of 19.7 — one cell did that.', ITEM, headsize=15, bodysize=12.5)
table(sl, 0.85, 4.3, [1.6, 2.2, 7.8],
      [['Mean-square', 'Pill', 'Reading'],
       ['0.7 – 1.5', ('green', GRN), 'Behaves as the model expects'],
       ['1.5 – 2.0', ('amber', ITEM), 'Noisy — check whether the item means the same thing to everyone'],
       ['above 2.0', ('red', MIS), 'Badly noisy — very likely measuring something else; look at this first'],
       ['below 0.7', ('amber', ITEM), 'Too predictable — often a near-duplicate; inflates reliability slightly'],
       ['thin evidence', ('grey', INK3), 'Provisional — under 5 on the minority side; a flagged fit is not a finding yet']],
      size=12, rowh=0.36)
txt(sl, 'Misfit is where the interesting work is. Do not delete a misfitting item without looking at who succeeded on it and who did not — that pattern is the finding.',
    0.85, 6.6, 11.6, 0.3, size=11, italic=True, color=INK2)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 8 — Extremes
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Honest behaviour 2 of 3')
title(sl, 'Extreme scores are named, not extrapolated', size=32)
txt(sl, 'Omar holds all twelve badges. He is somewhere above the hardest skill, and the data cannot say where.',
    0.85, 2.05, 11.6, 0.5, size=16, color=INK2)
card(sl, 0.85, 2.85, 5.6, 2.55, 'What Winsteps does',
     'Treats a perfect score as a fractional one — 11.7 of 12 by default (EXTRSC = 0.3) — and reports a '
     'measure with a standard error. Precise-looking, and arbitrary: change the setting and the person '
     'moves with no new data. The SE implies a sampling distribution around a quantity that has none.',
     INK3, headsize=15, bodysize=12.5)
card(sl, 6.85, 2.85, 5.6, 2.55, 'What this tool does',
     'Draws Omar at the top of the map in a hollow marker with no number. The table says 12 of 12 · '
     'maximum. He enters nothing numeric — not the mean, not the reliability, not the export. The same '
     'for a person at the bottom, and for a skill everybody or nobody passed.',
     ACC, headsize=15, bodysize=12.5)
banner(sl, 0.85, 5.6, 11.6, 1.0,
       'The pile-up stays visible. The fiction stays out.',
       'If a third of the group tops out, that ceiling is one of the most useful things a Wright map can show, '
       'and it is the most obvious thing on the page. What it does not get is a made-up measure.',
       GRN_LT, RGBColor(0x86, 0xef, 0xac), GRN)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 9 — The Likert case: setup
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Honest behaviour 3 of 3')
title(sl, 'The case a summed score can never show you', size=32)
txt(sl, 'A neighbourhood participation questionnaire. 140 residents, seven statements, each answered 1 to 5.',
    0.85, 2.05, 11.6, 0.5, size=16, color=INK2)
bullets(sl, ['I know who to call when something goes wrong',
             'People here look out for each other',
             'I have a say in what happens on my block',
             'I could get 10 neighbours to a meeting',
             'I have helped organise something locally',
             'Local decisions reflect what people here want',
             'I would take a problem to the council myself'],
        0.85, 2.75, 5.8, 3.6, size=13.5, gap=5, bullet='·')
card(sl, 6.85, 2.75, 5.6, 1.7, 'With more than two categories',
     'the tool switches to the Andrich rating-scale model and estimates one extra thing: where on the ruler '
     'each step of the scale is crossed. These are the thresholds, shared across all seven statements — '
     'which is exactly what a Likert form assumes.', ACC, headsize=14, bodysize=12)
card(sl, 6.85, 4.6, 5.6, 1.7, 'They should climb',
     'The step into "agree" should be crossed higher on the ruler than the step into "neutral". If it is '
     'not, respondents are not using the categories in the order they were written — and no total '
     'score can tell you that.', ITEM, headsize=14, bodysize=12)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 10 — The Likert case: result
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The rating scale card')
title(sl, 'Category 4 is crossed below category 3', size=32)
table(sl, 0.85, 2.05, [1.3, 1.1, 1.0, 1.5, 4.6],
      [['Category', 'Used', '%', 'Threshold', 'Reading'],
       ['1', '215', '21.9', 'baseline', '—'],
       ['2', '276', '28.2', '−1.81', 'Used as expected'],
       ['3', '95', '9.7', '+0.63', 'Used as expected'],
       [('4', MIS), '205', '20.9', ('−0.31', MIS), ('Disordered — sits below the one beneath it', MIS)],
       ['5', '189', '19.3', '+1.48', 'Used as expected']],
      size=12, rowh=0.36, numcols=(1, 2, 3))
# threshold picture: a short ruler with the four steps on it
rx, ry, rw = 0.95, 4.75, 8.8
conn(sl, rx, ry, rx + rw, ry, RULER, 1.5)
def RX(v): return rx + (v + 2.5) / 5.0 * rw
for t in range(-2, 3):
    conn(sl, RX(t), ry - 0.06, RX(t), ry + 0.06, RULER, 1)
    label(sl, ('+' if t > 0 else '') + str(t), RX(t) - 0.2, ry + 0.1, 0.4, 8, INK3, PP_ALIGN.CENTER)
for k, F, up in [(2, -1.81, True), (3, 0.63, True), (4, -0.31, False), (5, 1.48, True)]:
    col = MIS if not up else ITEM
    oval(sl, RX(F), ry, 0.16, col)
    label(sl, 'into ' + str(k), RX(F) - 0.4, ry - 0.42 if up else ry + 0.34, 0.8, 9, col, PP_ALIGN.CENTER)
conn(sl, RX(0.63), ry + 0.1, RX(-0.31), ry + 0.34, MIS, 1.0, MSO_LINE_DASH_STYLE.DASH)
label(sl, 'thresholds on the logit ruler — the step into 4 is crossed 0.94 logits below the step into 3',
      rx, ry + 0.62, 8, 8, INK3, bold=False)
card(sl, 9.35, 2.05, 3.1, 2.15, 'What a total shows',
     'Nothing. Every resident gets a number from 7 to 35. The distribution looks normal. Reliability looks '
     'fine. The report goes to the board.', INK3, headsize=13, bodysize=11)
card(sl, 9.35, 4.35, 3.1, 2.3, 'What the thresholds show',
     'Category 3 is used under 10% of the time and there is no point on the ruler where it is the most '
     'likely answer. Residents are not distinguishing "neutral" from "agree". Collapse 3 and 4, re-run.',
     MIS, headsize=13, bodysize=11)
txt(sl, 'Generated from a model with the middle thresholds deliberately reversed — a recovery, not a discovery. But it is the shape of the real finding.',
    0.85, 5.85, 8.3, 0.6, size=11.5, italic=True, color=INK2)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 11 — Two honest limits
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Limits, stated in the tool')
title(sl, 'Two things it will not pretend')
card(sl, 0.85, 2.15, 5.6, 2.7, 'Detecting a subtle disordering needs respondents',
     'An earlier version of this demo had 60 residents and a 0.2-logit reversal. The estimator reported the '
     'thresholds as ordered — the reversal was real but smaller than the sampling noise. That is the '
     'estimator being honest. At 1,500 respondents it recovers a 0.2-logit reversal reliably. Not cleverness; '
     'sample size.', ACC, headsize=15, bodysize=12.5)
card(sl, 6.85, 2.15, 5.6, 2.7, 'Spacing is approximate; sequence is reliable',
     'The estimation method (JMLE) spreads its estimates 10–20% wider than the truth with few items. '
     'Winsteps applies a correction factor; this tool does not. The disordering diagnostic depends only on '
     'the order of the thresholds, which the bias does not touch.', ITEM, headsize=15, bodysize=12.5)
banner(sl, 0.85, 5.1, 11.6, 1.3,
       'The estimator has been checked against known answers',
       'Parameter recovery on simulated data: correlation 0.995 dichotomous (400 × 12) and 0.999 for a '
       '4-category rating scale (500 × 10), RMS error about 0.2 logits. Threshold order recovered correctly '
       'in strongly disordered, mildly disordered, and ordered conditions. The bug that nearly shipped — '
       'margins that disagreed after an extreme was removed — is caught by a check that they must agree.',
       ACC_LT, ACC, ACC_DK)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 12 — Separation and strata
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Does it hold together?')
title(sl, 'How many groups can the data tell apart?', size=32)
for i, (v, lab) in enumerate([('0.84', 'person reliability'), ('2.32', 'separation'),
                              ('3.4', 'strata'), ('2.31', 'SD (logits)'), ('0.92', 'RMSE')]):
    x = 0.85 + i * 2.35
    txt(sl, v, x, 2.05, 2.2, 0.7, size=36, bold=True, color=ACC_DK)
    txt(sl, lab.upper(), x, 2.75, 2.2, 0.3, size=9.5, bold=True, color=INK3)
txt(sl, 'The badge demo: 23 measurable people, 12 skills.', 0.85, 3.15, 8, 0.3, size=11, italic=True, color=INK2)
card(sl, 0.85, 3.7, 3.7, 1.95, 'Reliability',
     'Share of the spread in measures that is real rather than error: (SD² − MSE) / SD². Same fact as '
     'separation, as a ratio.', ACC, headsize=14, bodysize=12)
card(sl, 4.8, 3.7, 3.7, 1.95, 'Separation',
     'Spread of measures divided by the typical error. 2.32 means the measures spread more than twice as '
     'wide as they are uncertain.', ACC, headsize=14, bodysize=12)
card(sl, 8.75, 3.7, 3.7, 1.95, 'Strata — the plain-language one',
     '(4G + 1) / 3. About three distinguishable groups of people. Enough to sort the ledger into levels. '
     'Not enough to rank people individually.', GRN, headsize=14, bodysize=12)
banner(sl, 0.85, 5.8, 11.6, 0.9,
       'Three strata is a sort, not a ranking',
       'Ranking individuals from an instrument with three strata is the standard misuse — the tool says so, in words, beside the number.',
       AMB_LT, AMB_ED, AMB_TX)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 13 — Why this matters for RCN
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Why now')
title(sl, 'The badge ledger is already the data')
for i, (h, b) in enumerate([
        ('Persons × skills, badged or not', 'The classic dichotomous design. No survey needed — the observations already exist, made by a mentor rather than reported by the person. Observed, not asserted.'),
        ('A measure can go on a control chart', 'A total of ordinal responses is not on an interval scale and must not be charted. A Rasch measure is. The path: badge ledger → Rasch measure → RCN Process Behavior Charts.'),
        ('Comparable across neighbourhoods', 'Measures are person-free and item-free. A construct measured with one set of items in Superior, Arizona and another in Whatcom County stays on one ruler if some items are shared.'),
        ('It can be wrong, and says so', 'Data fit the model or they do not. A summed score cannot fail — that is its weakness, not its strength.')]):
    x = 0.85 + (i % 2) * 5.95
    y = 2.15 + (i // 2) * 2.05
    card(sl, x, y, 5.65, 1.85, h, b, ACC if i < 3 else GRN, headsize=15, bodysize=12.5)
txt(sl, 'Lineage: Rasch → Wright and Linacre at Chicago → Bill Mahoney finding measures in Judy Hibbard\'s data, which became the Patient Activation Measure. The same method, on a badge ledger, asks how far a person is from running a neighbourhood.',
    0.85, 6.25, 11.6, 0.6, size=11.5, italic=True, color=INK2)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 14 — Vocabulary, and what is left out
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Scope')
title(sl, 'Seven words, and that is deliberate')
xx = 0.85
for w_ in ['measure', 'SE', 'infit', 'outfit', 'reliability', 'separation', 'strata']:
    wd = 0.32 + len(w_) * 0.115
    rect(sl, xx, 2.15, wd, 0.42, WHITE, BORDER)
    txt(sl, w_, xx, 2.18, wd, 0.36, size=14, color=INK, align=PP_ALIGN.CENTER)
    xx += wd + 0.14
txt(sl, 'Deliberately absent:', 0.85, 2.85, 6, 0.35, size=13, color=INK2)
xx, yy = 0.85, 3.25
for w_ in ['partial credit', 'many-facet Rasch', 'DIF', 'PCA of residuals', 'extreme-score extrapolation', 'bias correction', 'anchoring']:
    wd = 0.32 + len(w_) * 0.105
    if xx + wd > 12.45:
        xx, yy = 0.85, yy + 0.55
    rect(sl, xx, yy, wd, 0.42, SURF, BORDER)
    txt(sl, w_, xx, yy + 0.03, wd, 0.36, size=13, color=INK3, align=PP_ALIGN.CENTER)
    xx += wd + 0.14
banner(sl, 0.85, 4.35, 11.6, 1.1,
       'The first version should be the version a co-op board can read',
       'All of these are legitimate and some are important. Every addition is something the board has to '
       'learn. The rating-scale model was added because Likert forms need it and the disordering check is '
       'the single most valuable thing Rasch offers a survey.',
       ACC_LT, ACC, ACC_DK)
banner(sl, 0.85, 5.6, 11.6, 1.1,
       'When a real need appears, the estimator is ready to build on',
       'A badge that different mentors award at different standards is what many-facet Rasch is for, and it '
       'will probably be the first addition. It goes on top of an estimator that has been verified to '
       'recover known parameters.',
       AMB_LT, AMB_ED, AMB_TX)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 15 — Next
# ═════════════════════════════════════════════════════════════════════════════
sl = slide(ACC_DK)
rect(sl, 0, 0, 13.33, 7.5, ACC_DK)
rect(sl, 0.85, 1.5, 1.6, 0.06, ITEM)
txt(sl, 'What we would need from you', 0.85, 1.95, 11, 0.8, size=34, bold=True, color=WHITE)
for i, (h, b) in enumerate([
        ('A real badge ledger', 'One row per person, one column per badge, 1 where it is held. Twenty people '
                                'and ten badges is enough to see whether the skills reach the people.'),
        ('One questionnaire you already use', 'Raw responses, one row per respondent, one column per statement. '
                                              'The thresholds will say whether the scale works as written.'),
        ('The item you suspect', 'Every instrument has one that "never quite fit". Naming it before the run '
                                 'is the fastest test of whether the method earns its keep.'),
        ('A room', 'The Wright map is meant to be read by a group with no training. Whether that is true is '
                   'something only a room can tell us.')]):
    y = 2.85 + i * 1.0
    rect(sl, 0.85, y, 0.055, 0.95, ITEM)
    txt(sl, h, 1.15, y + 0.02, 4.0, 0.45, size=17, bold=True, color=WHITE)
    txt(sl, b, 5.3, y, 7.1, 0.95, size=13, color=PALE2, spacing=1.15)
txt(sl, 'The tool, an introduction, and a full user manual are ready to try now.',
    0.85, 6.95, 11, 0.35, size=13, italic=True, color=PALE)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'tools', 'rcn-rasch-intro.pptx')
out = os.path.normpath(out)
prs.save(out)
print('Wrote', out, '·', len(prs.slides._sldIdLst), 'slides')
