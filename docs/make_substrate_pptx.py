"""
Generate substrate/rcn-substrate-intro.pptx — the RCN Substrate intro deck.

Companion to substrate/substrate-intro.html; same argument, deck shape.
Regenerate after editing:  python3 docs/make_substrate_pptx.py
Requires: pip install python-pptx

ALWAYS RENDER BEFORE CALLING IT DONE. python-pptx overflows a text box
silently — the XML is valid, the shape reports its nominal size, and the words
run off the slide only when something lays them out. Open it (Keynote will
export a PDF from the command line) and look at every slide.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette — cyan for the substrate, matching the intro page ───────────────
CY     = RGBColor(0x0e, 0x74, 0x90)
CY_LT  = RGBColor(0xcf, 0xfa, 0xfe)
CY_BG  = RGBColor(0xec, 0xfe, 0xff)
CY_DK  = RGBColor(0x15, 0x5e, 0x75)
GRN    = RGBColor(0x16, 0x65, 0x34)
GRN_LT = RGBColor(0xdc, 0xfc, 0xe7)
RD     = RGBColor(0x99, 0x1b, 0x1b)
RD_LT  = RGBColor(0xfe, 0xe2, 0xe2)
AMB    = RGBColor(0x92, 0x40, 0x0e)
AMB_LT = RGBColor(0xff, 0xfb, 0xeb)
BLACK  = RGBColor(0x0f, 0x17, 0x2a)
GREY   = RGBColor(0x47, 0x55, 0x69)
LGREY  = RGBColor(0x94, 0xa3, 0xb8)
SLATE  = RGBColor(0xe2, 0xe8, 0xf0)
WHITE  = RGBColor(0xff, 0xff, 0xff)
OFF    = RGBColor(0xf8, 0xfa, 0xfc)
IN_BG, IN, IN_DK = CY_BG, CY, CY_DK   # so the shared primitives keep their defaults

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Primitives ───────────────────────────────────────────────────────────────

def rgb(hexstr):
    n = int(hexstr.lstrip('#'), 16)
    return RGBColor(n >> 16 & 255, n >> 8 & 255, n & 255)


def slide():
    return prs.slides.add_slide(BLANK)


def bg(sl, color):
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = color


def rect(sl, x, y, w, h, fill=WHITE, line=None, lw=1.5):
    sh = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def txt(sl, text, x, y, w, h, size=20, bold=False, italic=False,
        color=BLACK, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return tb


def txlines(sl, lines, x, y, w, h, sizes=None, bolds=None, colors=None,
            align=PP_ALIGN.LEFT, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(sizes[i] if isinstance(sizes, list) else (sizes or 18))
        r.font.bold = bolds[i] if isinstance(bolds, list) else (bolds or False)
        r.font.color.rgb = colors[i] if isinstance(colors, list) else (colors or BLACK)
        r.font.name = 'Calibri'
    return tb


def box_txt(sl, lines, x, y, w, h, fill=IN_BG, line=IN, lw=1.5,
            size=15, bold_first=True, text_color=BLACK,
            top_color=None, align=PP_ALIGN.CENTER):
    sh = rect(sl, x, y, w, h, fill=fill, line=line, lw=lw)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)
    for i, line_text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(size if i > 0 else size + 1)
        r.font.bold = (bold_first and i == 0)
        r.font.color.rgb = (top_color or text_color) if i == 0 else text_color
        r.font.name = 'Calibri'
    return sh


def label(sl, text, color=IN_DK):
    txt(sl, text, 0.45, 0.18, 8, 0.3, size=9, bold=True, color=color)


def bar(sl, color=IN, y=0.55, h=0.035):
    rect(sl, 0, y, 13.33, h, fill=color)


def title_txt(sl, text, color=BLACK, size=34, y=0.62):
    txt(sl, text, 0.45, y, 12.4, 1.3, size=size, bold=True, color=color)


def foot(sl, text):
    txt(sl, text, 0.45, 6.95, 12.4, 0.35, size=10, color=LGREY)




# ── 1 · Title ────────────────────────────────────────────────────────────────
s = slide()
bg(s, BLACK)
rect(s, 0, 0, 13.33, 0.14, fill=CY)
txt(s, "RCN TOOLSET  ·  THE SUBSTRATE", 0.9, 1.75, 11, 0.4, size=13, bold=True, color=CY_LT)
txt(s, "RCN Substrate", 0.9, 2.25, 11.5, 1.25, size=54, bold=True, color=WHITE)
txlines(s, [
    "One graph holding every layer at once.",
    "The tools stop being applications with their own files",
    "and become lenses on the same data.",
], 0.9, 3.7, 11.5, 1.6, sizes=22, colors=SLATE, spacing=8)
txt(s, "Neo4j 5.26.4 Enterprise  ·  projection layer on port 8768  ·  three databases, one set of queries",
    0.9, 6.3, 11.5, 0.4, size=13, italic=True, color=LGREY)

# ── 2 · The problem ──────────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "THE PROBLEM")
bar(s)
title_txt(s, "Every tool owned its own file")
txlines(s, [
    "The same concept lived in a graph JSON, a map, a timeline,",
    "and sixteen separate drawings.",
    "Nothing anywhere could say those were the same thing.",
], 0.55, 2.1, 12.2, 1.8, sizes=24, colors=GREY, spacing=12)
box_txt(s, [
    "Two answers to one question.",
    "Not a sync problem. A question of which copy is true — and no way to ask it.",
], 0.55, 4.4, 12.2, 1.4, fill=CY_BG, line=CY, size=17, top_color=CY_DK)
foot(s, "The substrate is the decision to stop having two answers.")

# ── 3 · The sentence ─────────────────────────────────────────────────────────
s = slide()
bg(s, BLACK)
rect(s, 0, 0, 13.33, 0.14, fill=CY)
txt(s, "THE RULE EVERYTHING FOLLOWS FROM", 0.9, 1.3, 11, 0.4, size=13, bold=True, color=CY_LT)
txlines(s, [
    "The substrate owns meaning.",
    "The file owns appearance.",
], 0.9, 2.0, 11.5, 1.7, sizes=40, bolds=True, colors=WHITE, spacing=6)
txt(s, "Neither pretends to own the other, and anything derivable is computed rather than stored.",
    0.9, 3.9, 11.5, 0.6, size=20, color=CY_LT)
txt(s, "Almost every design question answers itself from that sentence —",
    0.9, 5.0, 11.5, 0.45, size=17, color=SLATE)
txt(s, "and every bug found so far was a violation of it.",
    0.9, 5.5, 11.5, 0.5, size=19, bold=True, color=WHITE)

# ── 4 · Who owns what ────────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHO OWNS WHAT")
bar(s)
title_txt(s, "Three places a fact can live")
cols = [
    ("NEO4J", "Schema label, the states of a concept,\npolarity, relation family, sources",
     "Meaning. It has to merge across drawings.", CY_BG, CY, CY_DK),
    ("THE FILE", "Coordinates, colour, duplicate\nplacements, node ids",
     "Appearance. It belongs to one drawing.", WHITE, SLATE, GREY),
    ("NOWHERE", "Gold, family colour,\nthe fallback layout",
     "Derivable. Storing it makes a second answer that can drift.", GRN_LT, GRN, GRN),
]
for i, (h, what, why, fill, ln, tc) in enumerate(cols):
    box_txt(s, [h] + what.split("\n") + ["", why], 0.55 + i * 4.16, 2.15, 3.95, 2.9,
            fill=fill, line=ln, size=13, top_color=tc, align=PP_ALIGN.LEFT)
box_txt(s, [
    "The third column is the one people skip.",
    "gold is never stored — it is size(sources) > 1, evaluated fresh on every request. "
    "Store it and you have two answers again.",
], 0.55, 5.35, 12.23, 1.35, fill=WHITE, line=SLATE, size=14, top_color=BLACK)
foot(s, "Every bug in the pitfalls list is a fact that ended up in the wrong column.")

# ── 5 · What a lens is ───────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHAT A LENS IS")
bar(s)
title_txt(s, "One query, one lens, one envelope")
txt(s, "Every projection returns the Graph Tool's own schema, declared canonical so renderers "
       "read it with zero translation. No adapters.",
    0.55, 1.9, 12.2, 0.8, size=18, color=GREY)
box_txt(s, [
    "The renderers stay dumb.",
    "They draw shapes from data and know nothing about the domain. When a lens seems to need a "
    "custom renderer, the projection is usually wrong — fix the query, not the renderer.",
], 0.55, 2.85, 12.23, 1.5, fill=CY_BG, line=CY, size=16, top_color=CY_DK)
txlines(s, [
    "That rule paid for itself when the driver wheel was built.",
    "Its rings mean causal distance rather than naming — a completely different tree —",
    "and the Sunburst needed no new geometry at all.",
], 0.55, 4.65, 12.2, 1.5, sizes=[19, 17, 17], bolds=[True, False, False],
    colors=[BLACK, GREY, GREY], spacing=7)
foot(s, "A new lens costs a query, not an application.")

# ── 6 · The lenses ───────────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "THE LENSES")
bar(s)
title_txt(s, "Six readings of one graph")
lens = [
    ("schema", "the substrate describing itself,\nwith live counts"),
    ("causal", "the CLD reading. Its nodes are\nSTATES, not concepts"),
    ("structure", "the ERD reading. Deliberately\nunsigned"),
    ("gold", "who drew what, and where\ntwo hands met"),
    ("vocabulary", "family to concept to states,\nas a tree"),
    ("drivers", "causality unrolled from one state,\ncarrying the sign along each path"),
]
for i, (n, d) in enumerate(lens):
    x = 0.55 + (i % 3) * 4.16
    y = 2.1 + (i // 3) * 1.75
    box_txt(s, ["/projection/" + n] + d.split("\n"), x, y, 3.95, 1.5,
            fill=WHITE, line=SLATE, size=13, top_color=CY_DK, align=PP_ALIGN.LEFT)
box_txt(s, [
    "A subgraph is not stored. WHERE 'org' IN c.sources IS the drawing.",
    "The sixteen drawings and their union are the same rows read two ways — nothing is duplicated "
    "to make both readings work.",
], 0.55, 5.7, 12.23, 1.15, fill=CY_BG, line=CY, size=14, top_color=CY_DK)

# ── 7 · The claim ────────────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "THE CLAIM BEING TESTED")
bar(s)
title_txt(s, "Only the data grows")
txt(s, "Three databases, read by byte-identical Cypher. ?db= selects which graph, never which query.",
    0.55, 1.9, 12.2, 0.6, size=19, color=GREY)
rows = [
    ("neo4j", "6 concepts · 6 states · 7 edges", "the reference. Sources are PEOPLE", CY_BG, CY),
    ("composite26", "25 concepts · 25 states · 57 edges", "the signed CLD. Sources are topics", WHITE, SLATE),
    ("aspects16", "24 concepts · 25 states · 70 edges", "the 16 drawings. Sources are topics", WHITE, SLATE),
]
for i, (n, c, w, fill, ln) in enumerate(rows):
    y = 2.75 + i * 1.0
    rect(s, 0.55, y, 2.6, 0.8, fill=fill, line=ln, lw=1.5)
    txt(s, n, 0.55, y + 0.2, 2.6, 0.4, size=15, bold=True, color=CY_DK, align=PP_ALIGN.CENTER)
    txt(s, c, 3.35, y + 0.05, 4.6, 0.4, size=15, color=BLACK)
    txt(s, w, 3.35, y + 0.42, 9.3, 0.4, size=13, color=GREY)
box_txt(s, [
    "If a lens ever needs different queries for the bigger graph, the architecture was not "
    "proven at n=6 — and we would rather know that now than at n=260.",
], 0.55, 5.85, 12.23, 0.95, fill=AMB_LT, line=AMB, size=15, top_color=AMB)
foot(s, "The six-node reference is hand-written, small enough to check by eye, and exercises every layer at once.")

# ── 8 · The four levels ──────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "THE MODEL")
bar(s)
title_txt(s, "Four levels, and where an edge attaches")
for i, (n, d) in enumerate([
        ("FAMILY", "the eight. A node, generated from families.js"),
        ("CONCEPT", "the merge key — what two drawings agree on"),
        ("STATE", "a way the concept is measured. A NODE, because there can be several"),
        ("INSTANCE", "a thing that happened, with a date and a place")]):
    y = 1.95 + i * 0.72
    rect(s, 0.55, y, 1.75, 0.58, fill=CY_LT, line=CY, lw=1.2)
    txt(s, n, 0.55, y + 0.12, 1.75, 0.35, size=12, bold=True, color=CY_DK, align=PP_ALIGN.CENTER)
    txt(s, d, 2.5, y + 0.09, 10.3, 0.45, size=15, color=GREY)
box_txt(s, [
    "Causal edges join STATES. Structural edges join CONCEPTS.",
    "\u201cCoherence of PURPOSE raises Effectiveness of ORG\u201d is about measured quantities. "
    "\u201can Org exists for a Purpose\u201d is true however effective that org is. "
    "Push the second down to states and it says something nobody meant.",
], 0.55, 5.0, 12.23, 1.7, fill=CY_BG, line=CY, size=15, top_color=CY_DK)
foot(s, "A state is a node, not a property, because a concept can be measured several ways at once.")

# ── 9 · What it refuses ──────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHAT IT REFUSES TO DO")
bar(s)
title_txt(s, "The refusals are the design")
for i, (h, d) in enumerate([
        ("Store layout", "where a node sits belongs to a drawing"),
        ("Store what it can compute", "gold, family colour, a family\u2019s weight"),
        ("Infer what only a human knows", "nothing separates a person from a topic in a list of strings"),
        ("Guess a relation family", "17 edges have none. They are counted and left alone"),
        ("Let one save delete another\u2019s work", "retract, then re-assert. An empty save deletes nothing"),
        ("Be edited in the Browser", "families are generated copies. An edit there vanishes silently")]):
    y = 1.95 + i * 0.78
    txt(s, "\u2715", 0.55, y, 0.4, 0.4, size=17, bold=True, color=RD)
    txt(s, h, 1.05, y - 0.03, 5.0, 0.4, size=16, bold=True, color=BLACK)
    txt(s, d, 6.1, y - 0.01, 6.7, 0.45, size=14, color=GREY)
foot(s, "Every refusal exists because the alternative shipped once and looked like working code.")

# ── 10 · What it caught ──────────────────────────────────────────────────────
s = slide()
bg(s, OFF)
label(s, "WHAT THE RULE HAS CAUGHT", color=RD)
bar(s, color=RD)
title_txt(s, "Each of these looked like working code")
items = [
    ("A silent success that drew nothing", "coordinates left out; a clean load of 6 nodes, an empty canvas"),
    ("A diff nobody would read", "44 insertions for a one-character change"),
    ("A merge key that did not merge", "Active Goal vs ActiveGoal — one concept, two nodes, never gold"),
    ("A concept holding one meaning", "two states of Affect; string length decided which survived"),
    ("A save that would have deleted every causal edge", "three drawings previewed ZERO edges"),
    ("An invented self-loop", "merging two placements turned an edge between them into a loop"),
]
for i, (h, d) in enumerate(items):
    y = 1.95 + i * 0.78
    txt(s, h, 0.55, y - 0.03, 6.4, 0.42, size=15, bold=True, color=BLACK)
    txt(s, d, 7.05, y - 0.01, 5.75, 0.45, size=13, color=GREY)
box_txt(s, [
    "The last one never reached disk because all sixteen files were previewed before anything was written.",
], 0.55, 6.55, 12.23, 0.6, fill=RD_LT, line=RD, size=13, top_color=RD)

# ── 11 · Close ───────────────────────────────────────────────────────────────
s = slide()
bg(s, BLACK)
rect(s, 0, 0, 13.33, 0.14, fill=CY)
txt(s, "THE BET UNDERNEATH", 0.9, 1.4, 11, 0.4, size=13, bold=True, color=CY_LT)
txlines(s, [
    "A neighborhood cannot reason about itself",
    "while every tool holds a private copy",
    "of the same words.",
], 0.9, 2.0, 11.5, 2.2, sizes=30, bolds=True, colors=WHITE, spacing=8)
txt(s, "Put the meaning in one graph. Keep the appearance in the files people actually drew. "
       "Compute everything computable, and refuse to guess the rest.",
    0.9, 4.5, 11.5, 0.9, size=17, color=CY_LT)
txt(s, "Six lenses so far. The next one costs a query, not an application.",
    0.9, 5.5, 11.5, 0.5, size=17, italic=True, color=SLATE)
txt(s, "substrate/api.py  ·  port 8768  ·  Introduction \u00b7 User Manual \u00b7 README \u00b7 ROUND-TRIP",
    0.9, 6.3, 11.5, 0.4, size=13, color=LGREY)

# ── Save ─────────────────────────────────────────────────────────────────────
here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, '..', 'substrate', 'rcn-substrate-intro.pptx')
prs.save(out)
print('Wrote', os.path.normpath(out), '—', len(prs.slides._sldIdLst), 'slides')
