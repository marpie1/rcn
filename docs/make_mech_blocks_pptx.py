"""
Generate tools/rcn-mech-blocks-intro.pptx — the Mech Blocks intro deck.

Companion to tools/mech-blocks-intro.md, mech-blocks-manual.md and
mech-blocks-reference.md; same argument, deck shape. Regenerate after editing:
    python3 docs/make_mech_blocks_pptx.py
Requires: pip install python-pptx

AUDIENCE: Marc, NDC groups who will never type code, and Ward Cunningham.
Leads with what a Mech is, names the difficulty (typing, and errors that
only appear after running), then shows the three visual answers — needs and
makes, the three lamps, beginner mode — and is plain about the limit: green
lamps mean put together right, not a guaranteed result.

Every number is real: 61 handbook scripts from mech.fed.wiki, 1,433 moves and
5,535 checks from tools/test-mech-blocks.js (October 2026). Block names, needs
and makes come from the CATALOG in tools/mech-blocks.html.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ── Palette — the tool's block group colours ────────────────────────────────
CTRL   = RGBColor(0xb4, 0x53, 0x09)   # Control, brown-amber
WEB    = RGBColor(0x1d, 0x4e, 0xd8)   # Pages and neighbors, blue
SHOW   = RGBColor(0x15, 0x80, 0x3d)   # Show results, green
DATA   = RGBColor(0x0f, 0x76, 0x6e)   # Data, teal
SERV   = RGBColor(0x7c, 0x3a, 0xed)   # Server, purple
ACC_DK = RGBColor(0x1e, 0x29, 0x3b)
GRN    = RGBColor(0x15, 0x80, 0x3d)
GRN_LT = RGBColor(0xdc, 0xfc, 0xe7)
RED    = RGBColor(0xb9, 0x1c, 0x1c)
RED_LT = RGBColor(0xfe, 0xf2, 0xf2)
AMB    = RGBColor(0xb4, 0x53, 0x09)
AMB_LT = RGBColor(0xfe, 0xf9, 0xc3)
AMB_ED = RGBColor(0xfd, 0xe0, 0x47)
AMB_TX = RGBColor(0x85, 0x4d, 0x0e)
BLUE_LT= RGBColor(0xdb, 0xea, 0xfe)
INK    = RGBColor(0x0f, 0x17, 0x2a)
INK2   = RGBColor(0x47, 0x55, 0x69)
INK3   = RGBColor(0x94, 0xa3, 0xb8)
GREY   = RGBColor(0xd1, 0xd5, 0xdb)
SURF   = RGBColor(0xfb, 0xfa, 0xf6)
BORDER = RGBColor(0xe2, 0xe8, 0xf0)
WHITE  = RGBColor(0xff, 0xff, 0xff)
PALE   = RGBColor(0x9c, 0xa9, 0xbd)
PALE2  = RGBColor(0xcb, 0xd5, 0xe1)
MONO   = 'Menlo'

CREDIT = 'Marc Pierson and Claude Opus 5.5 · October 2026'

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Primitives (as in make_rcn_rasch_pptx.py) ───────────────────────────────

def slide(bgc=SURF):
    sl = prs.slides.add_slide(BLANK)
    f = sl.background.fill
    f.solid()
    f.fore_color.rgb = bgc
    return sl


def rect(sl, x, y, w, h, fill=WHITE, line=None, lw=1.25, dash=None, shape=1):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = dash
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def pill(sl, x, y, w, h, fill, line, lw=1.25, dash=None):
    sh = rect(sl, x, y, w, h, fill, line, lw, dash, shape=5)   # rounded rectangle
    sh.adjustments[0] = 0.5
    return sh


def conn(sl, x1, y1, x2, y2, color, lw=1.2, dash=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    return c


def arrow(sl, x1, y1, x2, y2, color, lw=2.0):
    """A straight connector with an arrowhead at its end."""
    from pptx.oxml.ns import qn
    from lxml import etree
    c = conn(sl, x1, y1, x2, y2, color, lw)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'triangle')
    return c


def txt(sl, text, x, y, w, h, size=20, bold=False, italic=False, color=INK,
        align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP, font='Calibri'):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
    return tb


def label(sl, text, x, y, w, size, color, align=PP_ALIGN.LEFT, bold=True, font='Calibri', h=0.24):
    """Tight label with zero insets, for annotating drawings."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return tb


def kicker(sl, text, color=WEB):
    txt(sl, text.upper(), 0.85, 0.55, 11, 0.32, size=11, bold=True, color=color)


def title(sl, text, y=0.95, size=32, color=INK, w=11.6):
    txt(sl, text, 0.85, y, w, 1.1, size=size, bold=True, color=color, spacing=0.95)


def footer(sl, n):
    txt(sl, 'RCN · Mech Blocks', 0.85, 6.95, 6, 0.3, size=9, color=INK3)
    txt(sl, str(n), 12.1, 6.95, 0.4, 0.3, size=9, color=INK3, align=PP_ALIGN.RIGHT)


# ── Text fitting: PowerPoint never shrinks text, so measure before placing ──
CHAR_EM = 0.505
BOTTOM = 6.75


def _lines(text, w_in, size_pt):
    cpl = max(8, int(w_in * 72 / (size_pt * CHAR_EM)))
    return sum(max(1, -(-len(p) // cpl)) for p in text.split('\n'))


def _text_h(text, w_in, size_pt, spacing=1.12):
    return _lines(text, w_in, size_pt) * size_pt * spacing * 1.02 / 72


def card(sl, x, y, w, h, head, body, accent=WEB, headsize=16, bodysize=13.5):
    tw = w - 0.5
    while bodysize > 9.0 and 0.72 + _text_h(body, tw, bodysize) + 0.12 > h:
        bodysize -= 0.5
    rect(sl, x, y, w, h, WHITE, BORDER)
    rect(sl, x, y, 0.06, h, accent)
    txt(sl, head, x + 0.28, y + 0.18, w - 0.5, 0.45, size=headsize, bold=True, color=INK)
    txt(sl, body, x + 0.28, y + 0.7, tw, h - 0.85, size=bodysize, color=INK2, spacing=1.1)


def banner(sl, x, y, w, h, head, body, fill, edge, headcolor, size=13):
    tw = w - 0.6
    need = lambda s: 0.58 + _text_h(body, tw, s) + 0.12
    while size > 9.0 and need(size) > h:
        size -= 0.5
    rect(sl, x, y, w, h, fill, edge)
    txt(sl, head, x + 0.3, y + 0.13, tw, 0.38, size=15, bold=True, color=headcolor)
    txt(sl, body, x + 0.3, y + 0.55, tw, h - 0.65, size=size, color=INK2, spacing=1.12)


# ── Drawing blocks the way the tool draws them ──────────────────────────────
FAM = {  # family icon and colour, as FAMILIES in the tool
    'neighborhood': ('🏘', WEB), 'page': ('🏘', WEB), 'info': ('🏘', WEB),
    'aspect': ('🔗', SHOW), 'marker': ('🔗', SHOW), 'items': ('☰', RGBColor(0xc2, 0x41, 0x0c)),
}


def chip(sl, x, y, kind, key, state='ok'):
    """A needs/makes label. Returns its width."""
    icon, col = FAM.get(key, ('◆', INK2))
    text = f'{kind} {icon} {key}'
    w = 0.32 + len(text) * 0.072
    if kind == 'makes':
        fill, line, tc = (WHITE, GREY, INK3) if state == 'stuck' else (col, col, WHITE)
    else:
        fill, line, tc = (AMB, AMB, WHITE) if state == 'missing' else (WHITE, col, col)
    pill(sl, x, y, w, 0.27, fill, line, 1.0, MSO_LINE_DASH_STYLE.DASH if state == 'stuck' else None)
    label(sl, text, x, y + 0.015, w, 10, tc, PP_ALIGN.CENTER, bold=state != 'stuck')
    return w


def block(sl, x, y, w, name, args, color, chips=(), h=0.5, shows=False, trouble=False):
    rect(sl, x, y, w, h, WHITE, color, 1.75)
    rect(sl, x, y, 0.1, h, color)
    label(sl, name, x + 0.22, y + h / 2 - 0.12, 1.4, 14, color)
    nx = x + 0.3 + len(name) * 0.135
    if args:
        label(sl, args, nx, y + h / 2 - 0.12, 1.5, 13, INK, bold=False)
        nx += 0.2 + len(args) * 0.1
    for kind, key, state in chips:
        nx += chip(sl, nx, y + h / 2 - 0.135, kind, key, state) + 0.08
    if shows:
        w2 = 1.35
        pill(sl, nx, y + h / 2 - 0.135, w2, 0.27, WHITE, GRN, 1.0)
        label(sl, '⇒ shows a result', nx, y + h / 2 - 0.12, w2, 10, GRN, PP_ALIGN.CENTER)
        nx += w2 + 0.08
    if trouble:
        rect(sl, nx, y + h / 2 - 0.13, 0.26, 0.26, WHITE, AMB, 1.25)
        label(sl, '✖', nx, y + h / 2 - 0.12, 0.26, 10, AMB, PP_ALIGN.CENTER)


def cblock(sl, x, y, w, inner_h, name, color, hat=True):
    """A block with a mouth: header, coloured spine, foot. Children go at x+0.4, y+0.55."""
    total = 0.5 + inner_h + 0.2
    rect(sl, x, y, w, total, WHITE, color, 1.75)
    rect(sl, x, y, 0.1, total, color)
    label(sl, name, x + 0.22, y + 0.13, 1.4, 14, color)
    rect(sl, x + 0.1, y + total - 0.16, w - 0.1, 0.16, RGBColor(0xf6, 0xe7, 0xd7))
    return total


def lamp(sl, x, y, text, on):
    w = 0.55 + len(text) * 0.095
    pill(sl, x, y, w, 0.36, WHITE, GRN if on else INK3, 1.75)
    label(sl, '●', x + 0.12, y + 0.06, 0.2, 13, GRN if on else INK3)
    label(sl, text, x + 0.36, y + 0.06, w - 0.4, 12.5, INK if on else INK2, bold=False)
    return w


def code(sl, text, x, y, w, h, size=17):
    rect(sl, x, y, w, h, WHITE, BORDER)
    txt(sl, text, x + 0.25, y + 0.18, w - 0.4, h - 0.3, size=size, font=MONO, color=INK, spacing=1.2)


n = 0
def nxt():
    global n
    n += 1
    return n


# ═════════════════════════════════════════════════════════════════════════════
# 1 — Title
# ═════════════════════════════════════════════════════════════════════════════
sl = slide(ACC_DK)
nxt()
rect(sl, 0.85, 1.55, 1.6, 0.06, CTRL)
txt(sl, 'Mech Blocks', 0.85, 1.85, 11, 1.0, size=48, bold=True, color=WHITE)
txt(sl, "Building small wiki programs by snapping blocks together,\non top of Ward Cunningham's Mech",
    0.85, 2.95, 11.5, 1.2, size=24, color=PALE2, spacing=1.1)
# a small stack as the emblem
cblock(sl, 8.6, 4.35, 3.9, 1.5, 'CLICK', CTRL)
for i, (nm, col) in enumerate([('NEIGHBORS', WEB), ('WALK', WEB), ('PREVIEW', SHOW)]):
    rect(sl, 9.0, 4.85 + i * 0.5, 3.3, 0.42, WHITE, col, 1.5)
    rect(sl, 9.0, 4.85 + i * 0.5, 0.08, 0.42, col)
    label(sl, nm, 9.2, 4.94 + i * 0.5, 2.0, 13, col)
txt(sl, 'An exploration. Nothing in Mech, on any wiki, or in the RCN tools in use is changed.',
    0.85, 4.6, 7.2, 0.9, size=15, italic=True, color=PALE, spacing=1.15)
txt(sl, CREDIT, 0.85, 6.6, 8, 0.4, size=14, bold=True, color=WHITE)

# ═════════════════════════════════════════════════════════════════════════════
# 2 — What a Mech is
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'What Mech is', WEB)
title(sl, 'A Mech is a few lines of capital words')
code(sl, 'CLICK\n NEIGHBORS fed.wiki\n WALK 10 steps\n PREVIEW graph', 0.85, 2.05, 4.3, 2.25, size=20)
rows = [('CLICK', 'puts a ▶ button on the page and waits for a person', CTRL),
        ('NEIGHBORS', 'gathers every page on the nearby wiki sites', WEB),
        ('WALK', 'follows links ten steps at a time and keeps small graphs', WEB),
        ('PREVIEW', 'opens those graphs as a page you can look at', SHOW)]
for i, (nm, what, col) in enumerate(rows):
    y = 2.05 + i * 0.58
    label(sl, nm, 5.6, y + 0.05, 1.6, 15, col)
    txt(sl, what, 7.25, y, 5.2, 0.55, size=15, color=INK2)
banner(sl, 0.85, 4.75, 11.6, 1.3, 'The blocks share one notebook',
       'Mech calls it "state". NEIGHBORS writes the neighborhood into it. WALK reads the neighborhood and '
       'writes graphs. PREVIEW reads the graphs. A line indented under another belongs to it: here, all three '
       'wait for CLICK. Put WALK before NEIGHBORS and WALK finds nothing to read.',
       BLUE_LT, WEB, WEB)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 3 — Why it is hard
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The difficulty', RED)
title(sl, 'Easy for Ward to type. Hard for nearly everyone else.')
for i, (h, b) in enumerate([
        ('Spaces carry the meaning', 'Which block a line belongs to is decided by how many spaces start it. '
                                     'One space too few and a line quietly belongs somewhere else, or nowhere.'),
        ('Every word must be exact', 'Thirty-one capital words, each spelled exactly, each with its own rules '
                                     'about what may follow it and what may be indented under it.'),
        ('Mistakes appear after you run', 'Click ▶ and only then: "WALK expects state.neighborhood, like from '
                                          'NEIGHBORS." Ward made these messages clear and uniform; they still '
                                          'arrive after the fact.')]):
    card(sl, 0.85 + i * 3.95, 2.1, 3.7, 2.4, h, b, accent=RED)
txt(sl, "Ward's own inspirations were programs for children, where pieces only fit where they belong.",
    0.85, 4.95, 11.6, 0.6, size=17, italic=True, color=INK2)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 4 — Three borrowings
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Where the idea comes from', CTRL)
title(sl, 'Three programs Ward named, one idea from each')
for i, (h, b, col) in enumerate([
        ('Scratch', 'Blocks with a mouth. CLICK is drawn as a C that holds other blocks, so an empty CLICK '
                    'is visibly unfinished instead of an error after running.', CTRL),
        ('Snap!', 'Blocks that pass data. Each block shows what it needs and what it makes, and a drop where '
                  'the need is missing is flagged while you drag.', WEB),
        ('Etoys', 'Pull pieces from what is on the screen. Drop a wiki page link on the script and it becomes '
                  'FROM site/slug, ready to hold the blocks that use the page.', SHOW)]):
    card(sl, 0.85 + i * 3.95, 2.1, 3.7, 2.25, h, b, accent=col)
banner(sl, 0.85, 4.7, 11.6, 1.15, "The checks are Ward's own words",
       "Every message comes from the trouble messages in Mech's code, applied before the script runs instead "
       "of after.", GRN_LT, GRN, GRN)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 5 — The one rule
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The one rule', WEB)
title(sl, "Ward's text stays the source. Blocks are a second view of it.")
code(sl, 'CLICK\n NEIGHBORS fed.wiki\n WALK 10 steps\n PREVIEW graph', 0.85, 2.15, 3.6, 1.95, size=16)
label(sl, 'what FedWiki stores', 0.85, 4.42, 3.6, 11, INK3, PP_ALIGN.CENTER, bold=False)
arrow(sl, 4.65, 2.85, 5.75, 2.85, INK2)
arrow(sl, 5.75, 3.35, 4.65, 3.35, INK2)
label(sl, 'every drag rewrites it', 4.55, 2.5, 1.3, 9, INK2, PP_ALIGN.CENTER, bold=False)
label(sl, 'typing rebuilds them', 4.55, 3.45, 1.3, 9, INK2, PP_ALIGN.CENTER, bold=False)
cblock(sl, 5.95, 2.15, 3.2, 1.45, 'CLICK', CTRL)
for i, (nm, col) in enumerate([('NEIGHBORS', WEB), ('WALK', WEB), ('PREVIEW', SHOW)]):
    rect(sl, 6.35, 2.65 + i * 0.47, 2.6, 0.4, WHITE, col, 1.5)
    rect(sl, 6.35, 2.65 + i * 0.47, 0.08, 0.4, col)
    label(sl, nm, 6.55, 2.73 + i * 0.47, 2.0, 12, col)
label(sl, 'what you drag', 5.95, 4.42, 3.2, 11, INK3, PP_ALIGN.CENTER, bold=False)
for i, (big, small) in enumerate([('61', 'handbook scripts come back byte for byte'),
                                  ('61', 'nest exactly as Ward\'s own interpreter nests them'),
                                  ('1,433', 'block moves keep every line and every block\'s shape'),
                                  ('5,535', 'automated checks pass')]):
    y = 2.05 + i * 0.62
    txt(sl, big, 9.5, y, 1.25, 0.6, size=26, bold=True, color=GRN, align=PP_ALIGN.RIGHT)
    txt(sl, small, 10.85, y + 0.06, 2.1, 0.6, size=11.5, color=INK2, spacing=1.0)
banner(sl, 0.85, 4.95, 11.6, 1.3, 'Why this rule matters',
       'Because only the text is stored, a script built from blocks runs anywhere Mech runs, and journal, fork, '
       'search and Ward\'s interpreter all keep working. Someone without Mech Blocks sees ordinary Mech. '
       'Nothing about the stored page changes.',
       BLUE_LT, WEB, WEB)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 6 — Needs and makes
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The first cue', SHOW)
title(sl, 'Every block says what it needs and what it makes')
cblock(sl, 0.85, 2.0, 7.4, 1.65, 'CLICK', CTRL)
block(sl, 1.25, 2.5, 6.8, 'NEIGHBORS', 'fed.wiki', WEB, [('makes', 'neighborhood', 'ok')], h=0.48)
block(sl, 1.25, 3.04, 6.8, 'WALK', '10 steps', WEB, [('needs', 'neighborhood', 'ok'), ('makes', 'aspect', 'ok')], h=0.48)
block(sl, 1.25, 3.58, 6.8, 'PREVIEW', 'graph', SHOW, [('needs', 'aspect', 'ok')], h=0.48, shows=True)
txt(sl, 'A block fits when a block above it makes the thing it needs, and the picture is the same.',
    0.85, 4.55, 7.4, 0.7, size=14, italic=True, color=INK2, spacing=1.1)
fams = [('🏘', 'sites and pages', 'neighborhood, page, info'), ('🔗', 'graphs', 'aspect, marker'),
        ('☰', 'lists of items', 'items'), ('📄', 'files and text', 'tsv, txt, csv, json…'),
        ('🌡', 'readings and counts', 'temperature, tick'), ('🖥', 'from the server', 'result, commons…'),
        ('✏️', 'the turtle drawing', 'turtle'), ('◆', 'made by CODE', 'anything else')]
label(sl, 'EIGHT FAMILIES, ONE PICTURE EACH', 8.7, 2.0, 4, 10, INK3)
for i, (ic, nm, keys) in enumerate(fams):
    y = 2.35 + i * 0.5
    label(sl, ic, 8.7, y, 0.4, 16, INK)
    label(sl, nm, 9.15, y - 0.06, 3.4, 13, INK)
    label(sl, keys, 9.15, y + 0.17, 3.4, 10, INK3, bold=False)
rect(sl, 0.85, 5.35, 7.4, 1.15, WHITE, BORDER)
block(sl, 1.05, 5.5, 7.0, 'WALK', '10 steps', WEB, [('needs', 'neighborhood', 'missing'), ('makes', 'aspect', 'stuck')], h=0.42, trouble=True)
txt(sl, 'Nothing above makes the neighborhood: the need turns amber, and "makes" is greyed out, because a block '
        'that stops early makes nothing.', 1.05, 5.95, 7.0, 0.55, size=11.5, color=INK2, spacing=1.05)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 7 — Three lamps
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The second cue', GRN)
title(sl, 'Three lamps answer three plain questions')
x = 0.85
for t in ['Fits together', 'Ready to start', 'Ends in a result']:
    x += lamp(sl, x, 2.0, t, True) + 0.15
for i, (h, b) in enumerate([
        ('Fits together', 'Every block has what it needs, every mouth that must be filled is filled, the words '
                          'after each block make sense, and no lines sit where they can never run.'),
        ('Ready to start', 'Anyone reading the page can start it. A CODE block not directly inside CLICK or TICK '
                           'runs only for the page owner: code travels with copied pages, and Mech waits for a '
                           'person to ask.'),
        ('Ends in a result', 'The script reaches a block that shows something (PREVIEW, REPORT, SOLO, POPUP, '
                             'PRINT, DOWNLOAD, SHOW or HELLO) and that block has what it needs.')]):
    card(sl, 0.85 + i * 3.95, 2.7, 3.7, 2.1, h, b, accent=GRN, bodysize=13)
banner(sl, 0.85, 5.1, 11.6, 1.2, 'What green cannot promise',
       'All three green means the script is put together right. A site may still not answer, a neighborhood '
       'may be empty, a sensor may be off. Only running shows those, and Mech reports them with its own ✖︎.',
       AMB_LT, AMB_ED, AMB_TX)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 8 — Beginner mode
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'The third cue', RED)
title(sl, 'Beginner mode: what does not fit cannot be dropped')
# palette with grey tiles
rect(sl, 0.85, 2.0, 2.6, 4.45, WHITE, INK, 1.5)
label(sl, 'Blocks', 1.05, 2.12, 2, 13, INK)
tiles = [('CLICK', CTRL, True), ('NEIGHBORS', WEB, True), ('WALK', WEB, False), ('RANDOM', WEB, False),
         ('FROM', WEB, True), ('PREVIEW', SHOW, False), ('SOLO', SHOW, False), ('HELLO', SHOW, True)]
for i, (nm, col, ok) in enumerate(tiles):
    y = 2.5 + i * 0.47
    c = col if ok else GREY
    rect(sl, 1.05, y, 2.2, 0.38, WHITE, c, 1.5)
    rect(sl, 1.05, y, 0.07, 0.38, c)
    label(sl, nm, 1.22, y + 0.07, 1.9, 12, col if ok else INK3)
label(sl, 'empty script: WALK, PREVIEW, SOLO are greyed', 0.85, 6.5, 3.6, 9.5, INK3, bold=False)
# refused drop
cblock(sl, 3.95, 2.0, 4.6, 1.15, 'CLICK', CTRL)
conn(sl, 4.35, 2.5, 8.35, 2.5, RED, 3.5)
block(sl, 4.35, 2.58, 4.0, 'NEIGHBORS', '', WEB, h=0.42)
block(sl, 4.35, 3.06, 4.0, 'WALK', '10 steps', WEB, h=0.42)
pill(sl, 4.35, 3.85, 4.2, 0.62, RED, RED)
txt(sl, 'Won\'t fit here: WALK expects "neighborhood", like from NEIGHBORS.', 4.45, 3.87, 4.0, 0.6,
    size=11, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
label(sl, 'dragging WALK above NEIGHBORS', 3.95, 4.55, 4.6, 9.5, INK3, PP_ALIGN.CENTER, bold=False)
txt(sl, '— Palette blocks that cannot go at the blue "next block goes here" line are greyed.\n'
        '— Tap a block to add it at the blue line: a script can be built by tapping alone.\n'
        '— A drop that would add a problem anywhere is refused, with the reason. Moving NEIGHBORS below '
        'WALK is refused too, because it breaks WALK.\n'
        '— Old problems do not block new drops, and an empty CLICK may be placed and filled later.\n'
        '— Turn it off and every drop goes through, amber, with the lamps and reasons still shown.',
    8.85, 2.0, 3.75, 4.5, size=12.5, color=INK2, spacing=1.12)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 9 — Four taps
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Tried in the browser', GRN)
title(sl, 'Four taps from an empty page to a working script')
steps = [('Tap CLICK', 'CLICK', 'CLICK expects indented blocks to follow.', (False, True, False)),
         ('Tap NEIGHBORS', 'CLICK\n NEIGHBORS', 'Nothing shows a result yet.', (True, True, False)),
         ('Tap WALK', 'CLICK\n NEIGHBORS\n WALK 10 steps', 'Nothing shows a result yet.', (True, True, False)),
         ('Tap PREVIEW', 'CLICK\n NEIGHBORS\n WALK 10 steps\n PREVIEW graph', '', (True, True, True))]
for i, (h, text, why, on) in enumerate(steps):
    x = 0.85 + i * 3.0
    label(sl, f'{i + 1}  {h}', x, 2.0, 2.8, 15, INK)
    code(sl, text, x, 2.4, 2.8, 1.75, size=13)
    for j, nm in enumerate(['Fits', 'Ready', 'Result']):
        pill(sl, x + j * 0.93, 4.35, 0.85, 0.32, WHITE, GRN if on[j] else INK3, 1.5)
        label(sl, '● ' + nm, x + j * 0.93, 4.39, 0.85, 10, GRN if on[j] else INK3, PP_ALIGN.CENTER)
    if why:
        txt(sl, why, x, 4.8, 2.8, 0.7, size=11, italic=True, color=AMB_TX, spacing=1.0)
banner(sl, 0.85, 5.55, 11.6, 0.95, 'Each tap lands where it belongs',
       'The blue line goes inside an empty CLICK, then follows the last block added. WALK only lit up in the '
       'palette once NEIGHBORS was there.', GRN_LT, GRN, GRN)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 10 — What reading Ward's code turned up
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, "Reading Ward's code", SERV)
title(sl, 'What the source showed that the handbook does not')
card(sl, 0.85, 2.0, 5.7, 2.15, 'HERE and THERE',
     'A short script on Ward\'s Catalog of Mech Blocks calls the blocks in the running code "here" and the '
     'blocks the page documents "there". The difference is the undocumented blocks: CODE and DOWNLOAD. '
     'Handbook pages for both are drafted, in Ward\'s template.', accent=SERV)
card(sl, 0.85, 4.35, 5.7, 2.15, 'A catalog Mech could carry',
     'What every block needs, makes and may hold now exists as a plain list, built by reading the code. '
     'Offered to Ward so tools like this need not work it out again.', accent=SERV)
for i, (h, b) in enumerate([
        ('Permission reaches one level', 'CLICK lets a CODE directly inside it run for anyone, but not one '
                                         'nested inside FROM. Ward\'s own Catalog of Variables runs only for him.'),
        ('FILE keeps the dot', '"FILE .txt" stores the text under ".txt", not "txt" as the handbook says, so KWIC '
                               'and DOWNLOAD cannot find it.'),
        ('Typography stops CODE', 'Curly quotes, dashes or emoji in a Code item stop it loading at all.')]):
    card(sl, 6.85, 2.0 + i * 1.5, 5.6, 1.38, h, b, accent=AMB, headsize=14, bodysize=12)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 11 — What it is and is not
# ═════════════════════════════════════════════════════════════════════════════
sl = slide()
kicker(sl, 'Honest scope', INK2)
title(sl, 'What it is, and what it is not')
card(sl, 0.85, 2.0, 5.7, 3.3, 'It is',
     '— One web page, tools/mech-blocks.html, that works offline.\n'
     '— A drag-and-tap editor for Mech text, with Ward\'s 61 handbook scripts to explore.\n'
     '— Checks before running, in Ward\'s words.\n'
     '— An introduction, a manual and a block reference, as FedWiki pages.\n'
     '— Tested: 5,535 automated checks, and real drags and taps in a browser.', accent=GRN, bodysize=14)
card(sl, 6.85, 2.0, 5.6, 3.3, 'It is not',
     '— It does not run Mech. Copy the text into a Mech item to run it.\n'
     '— It cannot see what CODE functions or PLUGIN blocks do.\n'
     '— It cannot know whether a site answers until the script runs.\n'
     '— Not yet a FedWiki plugin.\n'
     '— Not a change to anything in use: Mech, the wikis and the RCN tools are untouched.', accent=RED, bodysize=14)
footer(sl, nxt())

# ═════════════════════════════════════════════════════════════════════════════
# 12 — Next
# ═════════════════════════════════════════════════════════════════════════════
sl = slide(ACC_DK)
rect(sl, 0.85, 1.3, 1.6, 0.06, CTRL)
txt(sl, 'What could come next', 0.85, 1.6, 11, 0.8, size=34, bold=True, color=WHITE)
for i, (h, b) in enumerate([
        ('Talk with Ward', 'Bring the CODE and DOWNLOAD pages, the three findings and the catalog. Ask whether '
                           'he sees blocks as a view onto his text, or a replacement for it.'),
        ('Put it where people can try it', 'Post the page on NDC Assets so Ward and NDC groups can use it '
                                           'before any decision.'),
        ('Draw what Mech finds', 'WALK already makes graphs of the neighborhood. Hand them to the RCN Graph '
                                 'Tool, Map or Timeline: Mech finds, our tools draw.'),
        ('Then, a plugin', 'A block view that opens on any Mech item in FedWiki, once Ward has answered.')]):
    y = 2.65 + i * 0.95
    rect(sl, 0.85, y, 0.055, 0.85, CTRL)
    txt(sl, h, 1.15, y + 0.02, 4.0, 0.45, size=17, bold=True, color=WHITE)
    txt(sl, b, 5.3, y, 7.1, 0.9, size=13.5, color=PALE2, spacing=1.12)
txt(sl, CREDIT, 0.85, 6.65, 8, 0.4, size=13, bold=True, color=WHITE)
txt(sl, 'Mech is Ward Cunningham\'s work: github.com/WardCunningham/wiki-plugin-mech, handbook at mech.fed.wiki.',
    0.85, 7.0, 11.6, 0.35, size=11, italic=True, color=PALE)

out = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'tools', 'rcn-mech-blocks-intro.pptx'))
prs.save(out)
print('Wrote', out, '·', len(prs.slides._sldIdLst), 'slides')
